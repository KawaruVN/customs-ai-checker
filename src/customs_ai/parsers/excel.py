from datetime import datetime
from pathlib import Path
from typing import Any
import zipfile

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from customs_ai.parsers.errors import FileCorruptedError, ParserFailedError
from customs_ai.parsers.models import ParsedCell, ParsedSheet, ParsedWorkbook


class XlsxParser:
    parser_name = "openpyxl"

    def parse(self, document_id: str, file_path: Path) -> ParsedWorkbook:
        workbook = None
        try:
            workbook = load_workbook(
                filename=file_path,
                read_only=False,
                data_only=False,
                keep_vba=False,
                keep_links=False,
            )
            sheets: list[ParsedSheet] = []
            for sheet_index, worksheet in enumerate(workbook.worksheets):
                cells: list[ParsedCell] = []
                # openpyxl's worksheet dimensions can be inflated by formatting.  _cells
                # contains only instantiated cells and prevents scanning millions of empties.
                physical_cells = sorted(
                    (cell for cell in worksheet._cells.values() if isinstance(cell, Cell)),
                    key=lambda cell: (cell.row, cell.column),
                )
                for cell in physical_cells:
                    if cell.value is None:
                        continue
                    is_formula = cell.data_type == "f"
                    cells.append(
                        ParsedCell(
                            coordinate=cell.coordinate,
                            row=cell.row,
                            column=cell.column,
                            raw_value=cell.value,
                            displayed_value=None if is_formula else cell.value,
                            data_type=str(cell.data_type),
                            is_formula=is_formula,
                            number_format=cell.number_format,
                        )
                    )

                hidden_rows = sorted(
                    row_index
                    for row_index, dimension in worksheet.row_dimensions.items()
                    if dimension.hidden
                )
                hidden_columns = sorted(
                    column_key
                    for column_key, dimension in worksheet.column_dimensions.items()
                    if dimension.hidden
                )

                sheets.append(
                    ParsedSheet(
                        name=worksheet.title,
                        index=sheet_index,
                        state=worksheet.sheet_state,
                        merged_ranges=[str(value) for value in worksheet.merged_cells.ranges],
                        hidden_rows=hidden_rows,
                        hidden_columns=hidden_columns,
                        cells=cells,
                    )
                )

            return ParsedWorkbook(
                document_id=document_id,
                file_type="xlsx",
                parser_used=self.parser_name,
                workbook_name=file_path.name,
                sheet_count=len(sheets),
                sheets=sheets,
            )
        except (InvalidFileException, zipfile.BadZipFile, OSError, ValueError, KeyError, EOFError) as exc:
            raise FileCorruptedError("xlsx") from exc
        except FileCorruptedError:
            raise
        except Exception as exc:
            raise ParserFailedError() from exc
        finally:
            if workbook is not None:
                workbook.close()


class XlsParser:
    parser_name = "xlrd"

    def parse(self, document_id: str, file_path: Path) -> ParsedWorkbook:
        try:
            import xlrd
        except ImportError as exc:  # pragma: no cover - dependency contract guard
            raise ParserFailedError() from exc

        workbook = None
        try:
            workbook = xlrd.open_workbook(str(file_path), formatting_info=True, on_demand=True)
            sheets: list[ParsedSheet] = []
            for sheet_index in range(workbook.nsheets):
                worksheet = workbook.sheet_by_index(sheet_index)
                cells: list[ParsedCell] = []
                for row_index in range(worksheet.nrows):
                    row_length = worksheet.row_len(row_index)
                    for column_index in range(row_length):
                        cell = worksheet.cell(row_index, column_index)
                        if cell.ctype == xlrd.XL_CELL_EMPTY:
                            continue
                        raw_value: Any = cell.value
                        data_type = _xls_cell_type_name(xlrd, cell.ctype)
                        if cell.ctype == xlrd.XL_CELL_BOOLEAN:
                            raw_value = bool(cell.value)
                        elif cell.ctype == xlrd.XL_CELL_DATE:
                            try:
                                raw_value = xlrd.xldate_as_datetime(cell.value, workbook.datemode)
                            except Exception:
                                raw_value = cell.value

                        cells.append(
                            ParsedCell(
                                coordinate=f"{get_column_letter(column_index + 1)}{row_index + 1}",
                                row=row_index + 1,
                                column=column_index + 1,
                                raw_value=raw_value,
                                displayed_value=raw_value,
                                data_type=data_type,
                                is_formula=False,
                                number_format=None,
                            )
                        )

                merged_ranges = [
                    _xls_range_to_a1(xlrd, rlo, rhi, clo, chi)
                    for rlo, rhi, clo, chi in worksheet.merged_cells
                ]
                hidden_rows = sorted(
                    row_index + 1
                    for row_index, info in getattr(worksheet, "rowinfo_map", {}).items()
                    if getattr(info, "hidden", 0)
                )
                hidden_columns = sorted(
                    get_column_letter(column_index + 1)
                    for column_index, info in getattr(worksheet, "colinfo_map", {}).items()
                    if getattr(info, "hidden", 0)
                )
                visibility = getattr(worksheet, "visibility", 0)
                state = {0: "visible", 1: "hidden", 2: "veryHidden"}.get(visibility, "visible")
                sheets.append(
                    ParsedSheet(
                        name=worksheet.name,
                        index=sheet_index,
                        state=state,
                        merged_ranges=merged_ranges,
                        hidden_rows=hidden_rows,
                        hidden_columns=hidden_columns,
                        cells=cells,
                    )
                )

            return ParsedWorkbook(
                document_id=document_id,
                file_type="xls",
                parser_used=self.parser_name,
                workbook_name=file_path.name,
                sheet_count=len(sheets),
                sheets=sheets,
            )
        except (xlrd.biffh.XLRDError, OSError, ValueError, EOFError) as exc:
            raise FileCorruptedError("xls") from exc
        except FileCorruptedError:
            raise
        except Exception as exc:
            raise ParserFailedError() from exc
        finally:
            if workbook is not None:
                workbook.release_resources()


def _xls_cell_type_name(xlrd_module: Any, cell_type: int) -> str:
    names = {
        xlrd_module.XL_CELL_EMPTY: "empty",
        xlrd_module.XL_CELL_TEXT: "text",
        xlrd_module.XL_CELL_NUMBER: "number",
        xlrd_module.XL_CELL_DATE: "date",
        xlrd_module.XL_CELL_BOOLEAN: "boolean",
        xlrd_module.XL_CELL_ERROR: "error",
        xlrd_module.XL_CELL_BLANK: "blank",
    }
    return names.get(cell_type, "unknown")


def _xls_range_to_a1(xlrd_module: Any, row_low: int, row_high: int, col_low: int, col_high: int) -> str:
    start = xlrd_module.cellname(row_low, col_low)
    end = xlrd_module.cellname(row_high - 1, col_high - 1)
    return f"{start}:{end}"
