import hashlib
from datetime import datetime, timezone
from pathlib import Path

import pytest
from pypdf import PdfWriter

from customs_ai.application.parsing import DocumentParsingService
from customs_ai.config import settings
from customs_ai.ingestion.enums import ProcessingStatus
from customs_ai.ingestion.models import SourceDocument
from customs_ai.parsers.errors import (
    ParserFailedError,
    ParserUnsupportedError,
    SourceDocumentNotFoundError,
    SourceFileNotFoundError,
    SourceFilePathInvalidError,
)
from customs_ai.repositories.source_documents import SourceDocumentRepository


def _create_document(repo, document_id, file_type, path):
    document = SourceDocument(
        document_id=document_id,
        shipment_id="SHP-PARSE",
        original_filename=f"source.{file_type}",
        file_type=file_type,
        mime_type="application/octet-stream",
        file_hash=hashlib.sha256(path.read_bytes()).hexdigest() if path.exists() else "0" * 64,
        file_size=path.stat().st_size if path.exists() else 0,
        stored_path=str(path),
        upload_time=datetime.now(timezone.utc),
        processing_status=ProcessingStatus.UPLOADED,
    )
    repo.create(document)
    return document


def _blank_pdf(path):
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    with path.open("wb") as handle:
        writer.write(handle)


def test_missing_document_id_is_deterministic():
    service = DocumentParsingService(SourceDocumentRepository())
    with pytest.raises(SourceDocumentNotFoundError) as exc:
        service.parse_document("DOC-MISSING")
    assert exc.value.code == "SOURCE_DOCUMENT_NOT_FOUND"


def test_missing_source_file_is_deterministic():
    repo = SourceDocumentRepository()
    missing = settings.upload_root / "SHP" / "DOC" / "original.pdf"
    _create_document(repo, "DOC-MISSING-FILE", "pdf", missing)

    with pytest.raises(SourceFileNotFoundError) as exc:
        DocumentParsingService(repo).parse_document("DOC-MISSING-FILE")
    assert exc.value.code == "SOURCE_FILE_NOT_FOUND"


def test_path_outside_upload_root_is_rejected(tmp_path):
    repo = SourceDocumentRepository()
    outside = tmp_path / "outside.pdf"
    _blank_pdf(outside)
    _create_document(repo, "DOC-ESCAPE", "pdf", outside)

    with pytest.raises(SourceFilePathInvalidError) as exc:
        DocumentParsingService(repo).parse_document("DOC-ESCAPE")
    assert exc.value.code == "SOURCE_FILE_PATH_INVALID"


def test_unsupported_source_document_type_is_rejected():
    repo = SourceDocumentRepository()
    path = settings.upload_root / "SHP" / "DOC-CSV" / "original.csv"
    path.parent.mkdir(parents=True)
    path.write_text("a,b\n1,2", encoding="utf-8")
    _create_document(repo, "DOC-CSV", "csv", path)

    with pytest.raises(ParserUnsupportedError):
        DocumentParsingService(repo).parse_document("DOC-CSV")


def test_pdf_success_updates_parser_metadata_without_classification():
    repo = SourceDocumentRepository()
    path = settings.upload_root / "SHP" / "DOC-PDF" / "original.pdf"
    path.parent.mkdir(parents=True)
    _blank_pdf(path)
    before_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    _create_document(repo, "DOC-PDF", "pdf", path)

    parsed = DocumentParsingService(repo).parse_document("DOC-PDF")
    refreshed = repo.get_by_document_id("DOC-PDF")

    assert parsed.page_count == 1
    assert refreshed.processing_status == ProcessingStatus.PARSED
    assert refreshed.parser_used == "pypdf"
    assert refreshed.page_count == 1
    assert refreshed.sheet_count is None
    assert refreshed.detected_document_type is None
    assert refreshed.classification_confidence is None
    assert refreshed.extraction_status is None
    assert hashlib.sha256(path.read_bytes()).hexdigest() == before_hash


def test_repository_metadata_failure_does_not_mark_parsed_or_modify_source(monkeypatch):
    repo = SourceDocumentRepository()
    path = settings.upload_root / "SHP" / "DOC-FAIL" / "original.pdf"
    path.parent.mkdir(parents=True)
    _blank_pdf(path)
    before_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    _create_document(repo, "DOC-FAIL", "pdf", path)

    monkeypatch.setattr(repo, "update_parsing_metadata", lambda **kwargs: (_ for _ in ()).throw(RuntimeError("db down")))

    with pytest.raises(ParserFailedError):
        DocumentParsingService(repo).parse_document("DOC-FAIL")

    refreshed = repo.get_by_document_id("DOC-FAIL")
    assert refreshed.processing_status == ProcessingStatus.UPLOADED
    assert refreshed.parser_used is None
    assert hashlib.sha256(path.read_bytes()).hexdigest() == before_hash


def test_xlsx_success_updates_sheet_count_and_parser_used():
    from openpyxl import Workbook

    repo = SourceDocumentRepository()
    path = settings.upload_root / "SHP" / "DOC-XLSX" / "original.xlsx"
    path.parent.mkdir(parents=True)
    wb = Workbook()
    wb.active["A1"] = "Synthetic"
    wb.create_sheet("Second")["A1"] = 2
    wb.save(path)
    wb.close()
    _create_document(repo, "DOC-XLSX", "xlsx", path)

    parsed = DocumentParsingService(repo).parse_document("DOC-XLSX")
    refreshed = repo.get_by_document_id("DOC-XLSX")

    assert parsed.sheet_count == 2
    assert refreshed.processing_status == ProcessingStatus.PARSED
    assert refreshed.parser_used == "openpyxl"
    assert refreshed.sheet_count == 2
    assert refreshed.page_count is None
    assert refreshed.detected_document_type is None
    assert refreshed.classification_confidence is None
