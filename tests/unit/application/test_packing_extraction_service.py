"""Production Packing application flow, refusal boundaries and source safety."""

from contextlib import closing
from datetime import date, datetime, timezone
from decimal import Decimal
from hashlib import sha256
from types import SimpleNamespace

import openpyxl
import pytest
import xlwt
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from customs_ai.application.composition import build_packing_extraction_service
from customs_ai.application.packing_extraction import PackingExtractionService
from customs_ai.application.packing_models import PackingExtractionOutcome, PackingExtractionRequest
from customs_ai.application.parsing import DocumentParsingService
from customs_ai.classification.deterministic import DeterministicClassifier
from customs_ai.classification.config import load_classification_config
from customs_ai.classification.enums import DocumentType
from customs_ai.config import settings
from customs_ai.domain.schema import PackingList
from customs_ai.extraction.enums import ExtractionStatus
from customs_ai.extraction.packing_headers import DeterministicPackingHeaderExtractor
from customs_ai.extraction.packing_items import DeterministicPackingItemExtractor
from customs_ai.extraction.packing_models import PackingListHeaderExtractionResult, PackingItemExtractionResult
from customs_ai.ingestion.enums import ProcessingStatus
from customs_ai.ingestion.models import SourceDocument
from customs_ai.repositories.database import get_connection
from customs_ai.repositories.document_extractions import DocumentExtractionRepository
from customs_ai.repositories.source_documents import SourceDocumentRepository
from customs_ai.vision.enums import TextResolutionSource, VerificationLevel
from customs_ai.vision.models import ResolvedPdfDocument, ResolvedPdfPage


DOC="PACK-SERVICE"


def rows():
    return [["PACKING LIST"], ["PACKING LIST NO","000012"], ["PACKING LIST DATE","2026-10-29"],
            ["SELLER","Synthetic Seller"], ["BUYER","Synthetic Buyer"], ["INVOICE NO","000045"],
            ["INVOICE NO","000046"], ["TOTAL PACKAGES","0.00"],
            ["TOTAL GROSS WEIGHT","12345678901234567890.1200"], [],
            ["ITEM CODE","DESCRIPTION","QTY","UNIT","PACKAGE NO.","GROSS WEIGHT","NET WEIGHT"],
            ["000099","Synthetic 电机","0.00","PCS","000123","12.50","-2.50"]]


def workbook(*,file_type="xlsx", mutate=None, content=None):
    path=settings.upload_root/"synthetic"/("original."+file_type)
    path.parent.mkdir(parents=True,exist_ok=True)
    data=rows() if content is None else content
    if file_type=="xlsx":
        book=openpyxl.Workbook(); sheet=book.active; sheet.title="Synthetic Packing"
        for row in data: sheet.append(row)
        if mutate: mutate(sheet)
        book.save(path); book.close()
    else:
        book=xlwt.Workbook(); sheet=book.add_sheet("Synthetic Packing")
        for r, values in enumerate(data):
            for c,value in enumerate(values):
                if value is not None: sheet.write(r,c,value)
        book.save(str(path))
    return path


def pdf(blank=False):
    path=settings.upload_root/"synthetic"/"original.pdf"; path.parent.mkdir(parents=True,exist_ok=True)
    writer=PdfWriter(); page=writer.add_blank_page(width=400,height=400)
    if not blank:
        font=DictionaryObject({NameObject("/Type"):NameObject("/Font"),NameObject("/Subtype"):NameObject("/Type1"),NameObject("/BaseFont"):NameObject("/Helvetica")})
        page[NameObject("/Resources")]=DictionaryObject({NameObject("/Font"):DictionaryObject({NameObject("/F1"):font})})
        text=["PACKING LIST","PACKING LIST NO: 000012","PACKING LIST DATE: 2026-10-29","TOTAL GROSS WEIGHT: 12.50"]
        stream=DecodedStreamObject(); stream.set_data(("BT /F1 10 Tf 20 380 Td "+" 0 -15 Td ".join("("+line+") Tj" for line in text)+" ET").encode("ascii"))
        page[NameObject("/Contents")]=stream
    with path.open("wb") as handle: writer.write(handle)
    return path


def source(path, **updates):
    data=dict(document_id=DOC,shipment_id="SHP-PACKING",original_filename="misleading-invoice"+path.suffix,
        file_type=path.suffix[1:],mime_type="application/octet-stream",file_hash=sha256(path.read_bytes()).hexdigest(),
        file_size=path.stat().st_size,stored_path=str(path),upload_time=datetime.now(timezone.utc),
        processing_status=ProcessingStatus.CLASSIFIED,detected_document_type="PACKING_LIST",classification_confidence=0.9)
    data.update(updates); value=SourceDocument(**data); SourceDocumentRepository().create(value); return value


def service(**dependencies):
    return PackingExtractionService(SourceDocumentRepository(),**dependencies)


def history(): return DocumentExtractionRepository().list_history(DOC)


class ReadSpy:
    def __init__(self,callback=None): self.calls=0; self.states=[]; self.callback=callback
    def parse_document_readonly(self,document_id):
        self.calls+=1; repo=SourceDocumentRepository(); self.states.append(repo.get_by_document_id(document_id))
        result=DocumentParsingService(repo).parse_document_readonly(document_id)
        self.states.append(repo.get_by_document_id(document_id))
        if self.callback: self.callback(result)
        return result
    def parse_document(self,*args): raise AssertionError("Mutating parsing is forbidden")


class Vision:
    def __init__(self,mutate=None): self.calls=0; self.mutate=mutate
    def resolve_document(self,document_id,parsed):
        self.calls+=1
        result=ResolvedPdfDocument(document_id=document_id,page_count=1,visual_pages=[1],requires_review=False,
            pages=[ResolvedPdfPage(page=1,text="PACKING LIST\nPACKING LIST NO: 000012\nTOTAL GROSS WEIGHT: 12.50",
                source=TextResolutionSource.CROSS_VERIFIED_VISUAL,verification_level=VerificationLevel.CROSS_VERIFIED,
                agreement_score=0.9,confidence=0.95,requires_review=False)])
        if self.mutate: self.mutate(result)
        return result


@pytest.mark.parametrize("file_type",["xlsx","xls"])
def test_real_excel_pipeline_composes_all_values_v2_without_mutation(file_type):
    path=workbook(file_type=file_type); original=path.read_bytes(); src=source(path); reader=ReadSpy()
    outcome=service(parsing_service=reader).extract_document(DOC)
    assert outcome.persisted and outcome.result.status==ExtractionStatus.EXTRACTED
    assert outcome.record.is_active and outcome.active_record==outcome.record and reader.calls==1
    assert reader.states==[src,src]
    packing=outcome.result.packing_list
    assert packing.packing_list_number.normalized_value=="000012"
    assert packing.date.normalized_value==date(2026,10,29)
    assert [f.normalized_value for f in packing.invoice_references]==["000045","000046"]
    assert packing.package_count.normalized_value.as_tuple()==Decimal("0.00").as_tuple()
    assert packing.gross_weight.normalized_value.as_tuple()==Decimal("12345678901234567890.1200").as_tuple()
    assert packing.items[0].item_code.normalized_value=="000099"
    assert packing.items[0].description.normalized_value=="Synthetic 电机"
    assert packing.items[0].quantity.cell=="C12" and packing.items[0].package_number.normalized_value=="000123"
    assert all(c.field_path.startswith(("header.","items.")) for c in outcome.result.candidates)
    assert DocumentExtractionRepository().get_active(DOC).result==outcome.result
    refreshed=SourceDocumentRepository().get_by_document_id(DOC)
    assert refreshed.processing_status==ProcessingStatus.EXTRACTED
    assert refreshed.parser_used==src.parser_used and refreshed.page_count==src.page_count
    assert path.read_bytes()==original


def test_factory_uses_production_dependencies_and_no_network_for_excel(monkeypatch):
    monkeypatch.setattr(settings.vision,"enabled",False)
    source(workbook()); outcome=build_packing_extraction_service().extract_document(DOC)
    assert outcome.persisted and outcome.result.status==ExtractionStatus.EXTRACTED


def test_safe_native_pdf_bypasses_vision_and_pdf_item_extractor():
    source(pdf()); vision=Vision()
    class NoItems:
        def extract(self,*args): raise AssertionError("PDF items unsupported")
    outcome=service(vision_service=vision,item_extractor=NoItems()).extract_document(DOC)
    assert outcome.persisted and outcome.result.status==ExtractionStatus.NEEDS_REVIEW and vision.calls==0
    assert outcome.result.packing_list.items==[]
    assert outcome.result.packing_list.packing_list_number.page==1
    assert any(i.code=="PDF_PACKING_ITEMS_UNSUPPORTED" for i in outcome.result.issues)


def test_safe_resolved_pdf_uses_one_vision_pass_and_preserves_page():
    source(pdf(blank=True)); vision=Vision(); outcome=service(vision_service=vision).extract_document(DOC)
    assert outcome.persisted and outcome.result.status==ExtractionStatus.NEEDS_REVIEW and vision.calls==1
    assert outcome.result.packing_list.packing_list_number.page==1
    assert outcome.result.packing_list.items==[]


@pytest.mark.parametrize("patch", ["document_id","file_type","page_count","review","unresolved","fallback","visual_duplicates","visual_missing","page_review","unverified","conflict","confidence","agreement","page_id","empty"])
def test_unsafe_or_inconsistent_visual_evidence_cannot_produce_business_fields(patch):
    source(pdf(blank=True))
    def mutate(doc):
        if patch in {"document_id","file_type"}: setattr(doc,patch,"WRONG")
        elif patch=="page_count": doc.page_count=2
        elif patch=="review": doc.requires_review=True
        elif patch=="unresolved": doc.unresolved_pages=[1]
        elif patch=="fallback": doc.vision_fallback_recommended=True
        elif patch=="visual_duplicates": doc.visual_pages=[1,1]
        elif patch=="visual_missing": doc.visual_pages=[]
        elif patch=="page_review": doc.pages[0].requires_review=True
        elif patch=="unverified": doc.pages[0].verification_level=VerificationLevel.UNVERIFIED
        elif patch=="conflict": doc.pages[0].verification_level=VerificationLevel.CONFLICT
        elif patch=="confidence": doc.pages[0].confidence=0.1
        elif patch=="agreement": doc.pages[0].agreement_score=0.0
        elif patch=="page_id": doc.pages[0].page=2
        elif patch=="empty": doc.pages[0].text=""
    outcome=service(vision_service=Vision(mutate)).extract_document(DOC)
    assert outcome.persisted and outcome.result.status==ExtractionStatus.NEEDS_REVIEW
    assert outcome.result.packing_list==PackingList() and not outcome.result.candidates


@pytest.mark.parametrize("kind",[None,"COMMERCIAL_INVOICE","UNKNOWN","BILL_OF_LADING"])
def test_stored_wrong_classification_is_refused_without_parsing_or_history(kind):
    src=source(workbook(),detected_document_type=kind); spy=ReadSpy()
    outcome=service(parsing_service=spy).extract_document(DOC)
    assert not outcome.persisted and not history() and spy.calls==0
    assert SourceDocumentRepository().get_by_document_id(DOC)==src


@pytest.mark.parametrize("confidence",[None,0.1,float("nan"),float("inf")])
def test_unsafe_stored_confidence_never_authorizes_extraction(confidence):
    source(workbook(),classification_confidence=confidence)
    outcome=service().extract_document(DOC)
    assert not outcome.persisted and not history()


def test_current_invoice_content_is_refused_despite_stored_packing_classification():
    data=[["COMMERCIAL INVOICE"],["INVOICE NO","000012"],["TOTAL AMOUNT","25.00"]]
    src=source(workbook(content=data)); outcome=service().extract_document(DOC)
    assert not outcome.persisted and not history()
    assert SourceDocumentRepository().get_by_document_id(DOC)==src


@pytest.mark.parametrize("tamper",["id","type","confidence","evidence","clue","location"])
def test_classifier_stub_cannot_assert_ungrounded_high_confidence(tamper):
    source(workbook())
    class Classifier:
        def classify(self,context):
            output=DeterministicClassifier(load_classification_config()).classify(context)
            if tamper=="id": output.document_id="OTHER"
            elif tamper=="type": output.document_type=DocumentType.COMMERCIAL_INVOICE
            elif tamper=="confidence": output.confidence=0.99
            elif tamper=="evidence": output.evidence=[]
            elif tamper=="clue": output.evidence[0].clue="FABRICATED"
            elif tamper=="location": output.evidence[0].cell="ZZ999"
            return output
    outcome=service(classifier=Classifier()).extract_document(DOC)
    assert not outcome.persisted and not history()


@pytest.mark.parametrize("unsafe",["hidden_sheet","hidden_row","hidden_column","formula"])
def test_hidden_or_formula_classification_clues_do_not_authorize_wrong_document(unsafe):
    def mutate(sheet):
        if unsafe=="hidden_sheet":
            hidden=sheet.parent.create_sheet("Hidden"); hidden["A1"]="PACKING LIST"; hidden.sheet_state="hidden"
        else:
            sheet["H1"]="PACKING LIST"
            if unsafe=="hidden_row": sheet.row_dimensions[1].hidden=True
            elif unsafe=="hidden_column": sheet.column_dimensions["H"].hidden=True
            elif unsafe=="formula": sheet["H1"]='="PACKING LIST"'
    source(workbook(content=[["Unrelated source"]],mutate=mutate))
    outcome=service().extract_document(DOC)
    assert not outcome.persisted and not history()


@pytest.mark.parametrize("mutation",["hidden_row","hidden_group","formula"])
def test_excel_review_issues_and_safe_evidence_survive_composition(mutation):
    def mutate(sheet):
        if mutation=="hidden_row": sheet.row_dimensions[12].hidden=True
        elif mutation=="hidden_group": sheet.column_dimensions.group("B","D",hidden=True)
        else: sheet["C12"]="=1+1"
    source(workbook(mutate=mutate)); outcome=service().extract_document(DOC)
    assert outcome.persisted and outcome.result.status==ExtractionStatus.NEEDS_REVIEW
    assert any(i.is_review_required for i in outcome.result.issues)
    if mutation=="formula":
        assert outcome.result.packing_list.items[0].quantity is None
        assert any(c.raw_value=="=1+1" for c in outcome.result.candidates)


@pytest.mark.parametrize("invalid",[None,""," ",123," PACK-SERVICE"])
def test_invalid_request_does_not_touch_history(invalid):
    outcome=service().extract_document(invalid)
    assert not outcome.persisted and outcome.result.status==ExtractionStatus.FAILED


@pytest.mark.parametrize("change",["hash","size","path_escape","missing_file","state","stale_expected_hash","stale_expected_status"])
def test_preflight_source_guards_never_write(change,tmp_path):
    path=workbook(); updates={}
    if change=="state": updates["processing_status"]=ProcessingStatus.UPLOADED
    if change=="path_escape":
        outside=tmp_path/"outside.xlsx"; outside.write_bytes(path.read_bytes()); updates["stored_path"]=str(outside)
    src=source(path,**updates)
    if change=="hash": path.write_bytes(path.read_bytes()+b"tamper")
    elif change=="size":
        with closing(get_connection()) as connection,connection: connection.execute("UPDATE source_documents SET file_size=file_size+1")
    elif change=="missing_file": path.unlink()
    arguments={"expected_source_file_hash":"0"*64} if change=="stale_expected_hash" else {"expected_processing_status":ProcessingStatus.FAILED} if change=="stale_expected_status" else {}
    outcome=service().extract_document(DOC,**arguments)
    assert not outcome.persisted and not history()


def test_explicit_replacement_failed_execution_preserves_prior_active():
    source(workbook()); first=service().extract_document(DOC)
    assert first.persisted
    refused=service().extract_document(DOC)
    assert not refused.persisted and refused.active_record==first.record
    class Broken:
        def extract(self,*args): raise RuntimeError("PRIVATE SOURCE RAW")
    failed=service(header_extractor=Broken()).extract_document(DOC,expected_active_extraction_id=first.record.extraction_id)
    assert failed.persisted and failed.result.status==ExtractionStatus.FAILED and not failed.record.is_active
    assert failed.active_record==first.record and DocumentExtractionRepository().get_active(DOC)==first.record
    assert SourceDocumentRepository().get_by_document_id(DOC).processing_status==ProcessingStatus.FAILED
    assert len(history())==2 and "PRIVATE SOURCE RAW" not in str(failed)


@pytest.mark.parametrize("mutation",["source_hash","source_status","source_metadata"])
def test_source_changes_during_readonly_execution_refuse_commit(mutation):
    path=workbook(); source(path)
    def callback(parsed):
        if mutation=="source_hash": path.write_bytes(path.read_bytes()+b"tamper")
        else:
            with closing(get_connection()) as connection,connection:
                query="UPDATE source_documents SET processing_status='FAILED'" if mutation=="source_status" else "UPDATE source_documents SET original_filename='changed.xlsx'"
                connection.execute(query)
    outcome=service(parsing_service=ReadSpy(callback)).extract_document(DOC)
    assert not outcome.persisted and not history()


def test_persistence_exception_is_uncommitted_failure_without_retry(caplog):
    source(workbook())
    class BrokenRepo:
        calls=0
        def get_active(self,*args): return None
        def save_attempt(self,*args):
            self.calls+=1; raise RuntimeError("PRIVATE SQL PARAMETERS")
    repo=BrokenRepo(); outcome=service(extraction_repository=repo).extract_document(DOC)
    assert not outcome.persisted and outcome.result.status==ExtractionStatus.FAILED and repo.calls==1 and not history()
    assert "PRIVATE SQL PARAMETERS" not in str(outcome) and "PRIVATE SQL PARAMETERS" not in caplog.text


@pytest.mark.parametrize("component",["header","items"])
def test_wrong_component_contract_commits_technical_failure(component):
    source(workbook())
    wrong=SimpleNamespace(extract=lambda _:PackingList())
    outcome=service(**({"header_extractor":wrong} if component=="header" else {"item_extractor":wrong})).extract_document(DOC)
    assert outcome.persisted and outcome.result.status==ExtractionStatus.FAILED
    assert not outcome.record.is_active


def test_mutated_foreign_provenance_from_header_cannot_be_persisted_as_extracted():
    source(workbook())
    class Wrong:
        def extract(self,document):
            result=DeterministicPackingHeaderExtractor().extract(document)
            result.packing_list.packing_list_number.source_document_id="OTHER"
            return result
    outcome=service(header_extractor=Wrong()).extract_document(DOC)
    assert outcome.persisted and outcome.result.status==ExtractionStatus.FAILED
    assert outcome.result.packing_list==PackingList()


def test_valid_source_raw_with_fabricated_normalized_value_is_not_trusted():
    source(workbook())
    class Wrong:
        def extract(self,document):
            result=DeterministicPackingHeaderExtractor().extract(document)
            result.packing_list.package_count.normalized_value=Decimal("999")
            return result
    outcome=service(header_extractor=Wrong()).extract_document(DOC)
    assert outcome.persisted and outcome.result.status==ExtractionStatus.FAILED
    assert not outcome.record.is_active


def test_genuine_document_footer_is_retained_while_supported_goods_rows_are_routed():
    data=rows()+[[],["TOTAL NET WEIGHT","20.50"]]
    source(workbook(content=data)); outcome=service().extract_document(DOC)
    assert outcome.persisted and outcome.result.status==ExtractionStatus.EXTRACTED
    assert outcome.result.packing_list.net_weight.normalized_value==Decimal("20.50")
    assert outcome.result.packing_list.net_weight.cell=="B14"
    assert len(outcome.result.packing_list.items)==1
    assert any(c.field_path=="header.net_weight" and c.raw_value=="20.50" for c in outcome.result.candidates)


def test_contiguous_ambiguous_footer_review_from_item_component_remains_sticky():
    source(workbook(content=rows()+[["TOTAL NET WEIGHT","20.50"]]))
    outcome=service().extract_document(DOC)
    assert outcome.persisted and outcome.result.status==ExtractionStatus.NEEDS_REVIEW
    assert any(i.code=="AMBIGUOUS_PACKING_BOUNDARY" for i in outcome.result.issues)
    assert any(c.raw_value=="TOTAL NET WEIGHT" for c in outcome.result.candidates)


def test_mixed_document_label_outside_item_columns_cannot_silently_be_discarded():
    def mutate(sheet): sheet["M12"]="SELLER: Product annotation"
    source(workbook(mutate=mutate)); outcome=service().extract_document(DOC)
    assert outcome.persisted and outcome.result.status==ExtractionStatus.NEEDS_REVIEW
    assert outcome.result.packing_list.seller.name.normalized_value=="Synthetic Seller"


def test_injected_classifier_with_malformed_result_is_refused_not_committed():
    source(workbook()); outcome=service(classifier=SimpleNamespace(classify=lambda _:object())).extract_document(DOC)
    assert not outcome.persisted and not history()


def test_missing_source_is_uncommitted_failure():
    outcome=service().extract_document("ABSENT")
    assert not outcome.persisted and outcome.result.status==ExtractionStatus.FAILED


def test_source_repository_cannot_return_another_document_for_requested_identity():
    original=source(workbook())
    other=original.model_copy(update={"document_id":"OTHER"})
    outcome=PackingExtractionService(SimpleNamespace(get_by_document_id=lambda _:other)).extract_document(DOC)
    assert not outcome.persisted and not history()


def test_stale_active_request_is_refused_without_new_attempt():
    source(workbook()); first=service().extract_document(DOC)
    outcome=service().extract_document(DOC,expected_active_extraction_id="STALE")
    assert not outcome.persisted and outcome.active_record==first.record
    assert history()==[first.record]


def test_native_excel_date_type_survives_composition_and_sqlite():
    def mutate(sheet): sheet["B3"]=datetime(2026,10,29,13,45)
    source(workbook(mutate=mutate)); outcome=service().extract_document(DOC)
    assert outcome.persisted and outcome.result.status==ExtractionStatus.EXTRACTED
    assert type(outcome.result.packing_list.date.normalized_value) is date
    assert outcome.result.packing_list.date.raw_value=="2026-10-29 13:45:00"


def test_mutated_active_record_from_dependency_fails_closed_without_outcome_crash():
    source(workbook()); first=service().extract_document(DOC)
    wrong=first.record.model_copy(update={"document_id":"OTHER"})
    dependency=SimpleNamespace(get_active=lambda _:wrong)
    outcome=service(extraction_repository=dependency).extract_document(DOC)
    assert not outcome.persisted and outcome.result.status==ExtractionStatus.FAILED
    assert history()==[first.record]


@pytest.mark.parametrize("format",["xlsx","pdf"])
def test_substring_of_identifier_is_not_exact_source_evidence(format):
    source(workbook() if format=="xlsx" else pdf())
    class Wrong:
        def extract(self,document):
            result=DeterministicPackingHeaderExtractor().extract(document)
            result.packing_list.packing_list_number.raw_value="0"
            result.packing_list.packing_list_number.normalized_value="0"
            return result
    outcome=service(header_extractor=Wrong()).extract_document(DOC)
    assert outcome.persisted and outcome.result.status==ExtractionStatus.FAILED
    assert not outcome.record.is_active
