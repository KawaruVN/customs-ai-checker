from datetime import datetime, timezone
from hashlib import sha256
from types import SimpleNamespace

import pytest
from openpyxl import Workbook
from pypdf import PdfWriter

from customs_ai.application.parsing import DocumentParsingService
from customs_ai.config import settings
from customs_ai.ingestion.enums import ProcessingStatus
from customs_ai.ingestion.models import SourceDocument
from customs_ai.parsers.errors import (
    ParserFailedError,
    ParserUnsupportedError,
    SourceDocumentNotFoundError,
    SourceFileNotFoundError,
    SourceFilePathInvalidError,
)
from customs_ai.parsers.models import ParsedPdfDocument, ParsedPdfPage
from customs_ai.repositories.source_documents import SourceDocumentRepository


def _source(file_type="pdf", status=ProcessingStatus.CLASSIFIED, path=None):
    path = path or settings.upload_root / "synthetic" / f"original.{file_type}"
    path.parent.mkdir(parents=True, exist_ok=True)
    if file_type == "pdf":
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        with path.open("wb") as handle:
            writer.write(handle)
    elif file_type == "xlsx":
        book = Workbook()
        book.active["C8"] = "Synthetic source"
        book.save(path)
        book.close()
    else:
        path.write_bytes(b"synthetic")
    source = SourceDocument(
        document_id="DOC-READONLY", shipment_id="SHP-SYNTHETIC",
        original_filename=f"synthetic.{file_type}", file_type=file_type,
        mime_type="application/octet-stream", file_hash=sha256(path.read_bytes()).hexdigest(),
        file_size=path.stat().st_size, stored_path=str(path),
        upload_time=datetime.now(timezone.utc), processing_status=status,
        detected_document_type="COMMERCIAL_INVOICE", classification_confidence=0.95,
        parser_used="synthetic_metadata", page_count=17,
        extraction_status=status.value if status != ProcessingStatus.CLASSIFIED else None,
    )
    repo = SourceDocumentRepository()
    repo.create(source)
    return repo, source, path


@pytest.mark.parametrize("status", [ProcessingStatus.CLASSIFIED, ProcessingStatus.EXTRACTED, ProcessingStatus.NEEDS_REVIEW])
@pytest.mark.parametrize("file_type", ["pdf", "xlsx"])
def test_readonly_parsing_preserves_entire_source_metadata_and_bytes(status, file_type):
    repo, source, path = _source(file_type, status)
    before = path.read_bytes()
    result = DocumentParsingService(repo).parse_document_readonly(source.document_id)
    assert result.document_id == source.document_id
    assert result.file_type == file_type
    assert repo.get_by_document_id(source.document_id) == source
    assert path.read_bytes() == before
    assert sha256(path.read_bytes()).hexdigest() == source.file_hash


def test_readonly_path_never_calls_metadata_update(monkeypatch):
    repo, source, _ = _source()
    def forbidden(**kwargs):
        raise AssertionError("A readonly parse must never call a metadata writer")
    monkeypatch.setattr(repo, "update_parsing_metadata", forbidden)
    DocumentParsingService(repo).parse_document_readonly(source.document_id)
    assert repo.get_by_document_id(source.document_id) == source


def test_readonly_missing_document_retains_typed_failure():
    with pytest.raises(SourceDocumentNotFoundError):
        DocumentParsingService(SourceDocumentRepository()).parse_document_readonly("DOC-MISSING")


def test_readonly_missing_file_retains_typed_failure_and_metadata():
    repo, source, path = _source()
    path.unlink()
    with pytest.raises(SourceFileNotFoundError):
        DocumentParsingService(repo).parse_document_readonly(source.document_id)
    assert repo.get_by_document_id(source.document_id) == source


def test_readonly_source_path_containment_still_enforced(tmp_path):
    repo, source, _ = _source(path=tmp_path / "outside.pdf")
    with pytest.raises(SourceFilePathInvalidError) as failure:
        DocumentParsingService(repo).parse_document_readonly(source.document_id)
    assert str(tmp_path) not in str(failure.value)
    assert failure.value.__suppress_context__ is True
    assert repo.get_by_document_id(source.document_id) == source


def test_readonly_unsupported_type_retains_typed_failure():
    repo, source, _ = _source("csv")
    with pytest.raises(ParserUnsupportedError):
        DocumentParsingService(repo).parse_document_readonly(source.document_id)
    assert repo.get_by_document_id(source.document_id) == source


def _parsed():
    return ParsedPdfDocument(
        document_id="DOC-READONLY", page_count=1, has_text_layer=True,
        needs_ocr=False, pages=[ParsedPdfPage(page=1, text="Synthetic source")],
    )


@pytest.mark.parametrize("damage", ["wrong_document", "nested_invalid", "unsupported", "wrong_file_type"])
def test_readonly_parser_result_is_revalidated_before_return(damage):
    repo, source, _ = _source()
    parsed = _parsed()
    if damage == "wrong_document":
        parsed.document_id = "DOC-OTHER"
    elif damage == "nested_invalid":
        parsed.pages[0].text = {"unexpected": "synthetic"}
    elif damage == "wrong_file_type":
        parsed.file_type = "xlsx"
    else:
        parsed = object()
    router = SimpleNamespace(get_parser=lambda _: SimpleNamespace(parse=lambda *_: parsed))
    with pytest.raises(ParserFailedError):
        DocumentParsingService(repo, router).parse_document_readonly(source.document_id)
    assert repo.get_by_document_id(source.document_id) == source


def test_readonly_unexpected_parser_failure_is_sanitized(caplog):
    repo, source, path = _source()
    private = "SYNTHETIC_PRIVATE_CUSTOMER_VALUE"
    def fail(*args):
        raise RuntimeError(f"{private} {path}")
    router = SimpleNamespace(get_parser=lambda _: SimpleNamespace(parse=fail))
    with pytest.raises(ParserFailedError) as failure:
        DocumentParsingService(repo, router).parse_document_readonly(source.document_id)
    assert private not in str(failure.value)
    assert str(path) not in str(failure.value)
    assert failure.value.__cause__ is None
    assert failure.value.__suppress_context__ is True
    assert private not in caplog.text
    assert repo.get_by_document_id(source.document_id) == source


@pytest.mark.parametrize("document_id", [None, 42, "", " ", " DOC-READONLY"])
def test_readonly_malformed_identifier_is_controlled(document_id):
    with pytest.raises(ParserFailedError):
        DocumentParsingService(SourceDocumentRepository()).parse_document_readonly(document_id)
