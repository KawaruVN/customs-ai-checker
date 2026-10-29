"""Invoice application contract tests using synthetic sources and real SQLite."""

from contextlib import closing
from datetime import date, datetime, timezone
from decimal import Decimal
from hashlib import sha256

import openpyxl
import pytest
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from customs_ai.application.composition import build_invoice_extraction_service
from customs_ai.application.invoice_extraction import InvoiceExtractionService
from customs_ai.application.parsing import DocumentParsingService
from customs_ai.config import settings
from customs_ai.extraction.enums import ExtractionStatus
from customs_ai.extraction.headers import DeterministicInvoiceHeaderExtractor
from customs_ai.extraction.models import ExtractionCandidate
from customs_ai.extraction.semantic import SemanticInvoiceResponse
from customs_ai.ingestion.enums import ProcessingStatus
from customs_ai.ingestion.models import SourceDocument
from customs_ai.parsers.errors import ParserFailedError
from customs_ai.repositories.database import get_connection
from customs_ai.repositories.document_extractions import DocumentExtractionRepository
from customs_ai.repositories.source_documents import SourceDocumentRepository
from customs_ai.vision.enums import TextResolutionSource, VerificationLevel
from customs_ai.vision.models import ResolvedPdfDocument, ResolvedPdfPage


DOC = "DOC-SYNTHETIC-INVOICE"


def _path(suffix):
    path = settings.upload_root / "SHP-SYNTHETIC" / DOC / ("original." + suffix)
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def _xlsx(*, title="COMMERCIAL INVOICE", formula=False, conflict=False,
          items=True, incomplete=False, extra_table=False, freight=False):
    path = _path("xlsx")
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Synthetic Invoice"
    sheet["A1"] = title
    if "COMMERCIAL INVOICE" in title and "PACKING LIST" in title:
        # Balance the configured supporting clues so this is a genuine UNKNOWN
        # classification under the approved conflict-margin contract.
        sheet["A13"] = "GROSS WEIGHT"
        sheet["A14"] = "NET WEIGHT"
    for row, label, value in [
        (3, "INVOICE NO", "000012"),
        (4, "INVOICE DATE", date(2026, 10, 29)),
        (5, "SELLER", "Synthetic Seller"),
        (6, "BUYER", "Synthetic Buyer"),
        (7, "SHIP TO", "Synthetic Receiver"),
        (8, "CURRENCY", "USD"),
        (9, "TOTAL AMOUNT", "=G18" if formula else "25.00"),
    ]:
        sheet.cell(row, 1, label)
        sheet.cell(row, 2, value)
    if conflict:
        sheet["A11"] = "INVOICE NO"
        sheet["B11"] = "000099"
    if freight:
        sheet["A12"] = "FREIGHT"
        sheet["B12"] = "2.50"
    if items:
        for column, value in enumerate(["ITEM CODE", "DESCRIPTION", "QTY", "UNIT PRICE", "AMOUNT"], 3):
            sheet.cell(17, column, value)
        for column, value in enumerate(["000042", "Synthetic Motor", 2, 12.5, 25], 3):
            sheet.cell(18, column, value)
        if incomplete:
            sheet["D19"] = "Synthetic partial product"
        if extra_table:
            for column, value in enumerate(["ITEM CODE", "QTY"], 3):
                sheet.cell(25, column, value)
            sheet["C26"] = "000043"
            sheet["D26"] = 1
    workbook.save(path)
    workbook.close()
    return path


def _pdf(*, blank=False, pages=1):
    path = _path("pdf")
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    if not blank:
        font = DictionaryObject({
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        })
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})
        })
        lines = ["COMMERCIAL INVOICE", "INVOICE NO: 000012", "INVOICE DATE: 2026-10-29",
                 "SELLER: Synthetic Seller", "BUYER: Synthetic Buyer", "TOTAL AMOUNT: 25.00"]
        stream = DecodedStreamObject()
        text = "BT /F1 10 Tf 10 280 Td " + " 0 -15 Td ".join("(" + line + ") Tj" for line in lines) + " ET"
        stream.set_data(text.encode("ascii"))
        page[NameObject("/Contents")] = stream
    for _ in range(pages - 1):
        writer.add_blank_page(width=300, height=300)
    with path.open("wb") as handle:
        writer.write(handle)
    return path


def _source(path, **overrides):
    values = dict(
        document_id=DOC, shipment_id="SHP-SYNTHETIC", original_filename="misleading_invoice." + path.suffix[1:],
        file_type=path.suffix[1:], mime_type="application/octet-stream",
        file_hash=sha256(path.read_bytes()).hexdigest(), file_size=path.stat().st_size,
        stored_path=str(path), upload_time=datetime.now(timezone.utc),
        processing_status=ProcessingStatus.CLASSIFIED,
        detected_document_type="COMMERCIAL_INVOICE", classification_confidence=0.9,
        parser_used="previous_parser", page_count=7, sheet_count=8,
    )
    values.update(overrides)
    source = SourceDocument(**values)
    SourceDocumentRepository().create(source)
    return source


class SemanticSpy:
    def __init__(self, respond=None):
        self.requests = []
        self.respond = respond

    def extract(self, request):
        self.requests.append(request)
        return self.respond(request) if self.respond else SemanticInvoiceResponse(document_id=request.document_id, candidates=[])


class ReadSpy:
    def __init__(self, callback=None):
        self.calls = 0
        self.observations = []
        self.callback = callback

    def parse_document_readonly(self, document_id):
        repository = SourceDocumentRepository()
        self.calls += 1
        self.observations.append(repository.get_by_document_id(document_id))
        parsed = DocumentParsingService(repository).parse_document_readonly(document_id)
        self.observations.append(repository.get_by_document_id(document_id))
        if self.callback:
            self.callback(parsed)
        return parsed

    def parse_document(self, document_id):
        raise AssertionError("Extraction must not use the mutating parser API")


def _history():
    return DocumentExtractionRepository().list_history(DOC)


def _service(**dependencies):
    return InvoiceExtractionService(SourceDocumentRepository(), **dependencies)


def test_real_excel_composes_headers_items_and_persists_exact_canonical_provenance():
    path = _xlsx()
    source = _source(path)
    before = path.read_bytes()
    outcome = _service().extract_document(DOC)
    assert outcome.persisted
    assert outcome.result.status == ExtractionStatus.EXTRACTED
    invoice = outcome.result.invoice
    assert invoice.invoice_number.normalized_value == "000012"
    assert invoice.invoice_number.raw_value == "000012"
    assert invoice.invoice_number.source_document_id == DOC
    assert invoice.invoice_number.sheet == "Synthetic Invoice"
    assert invoice.invoice_number.cell == "B3"
    assert type(invoice.invoice_date.normalized_value) is date
    assert invoice.invoice_date.normalized_value == date(2026, 10, 29)
    assert invoice.seller.name.normalized_value == "Synthetic Seller"
    assert invoice.buyer.name.normalized_value == "Synthetic Buyer"
    assert invoice.ship_to.name.normalized_value == "Synthetic Receiver"
    assert invoice.total_amount.normalized_value == Decimal("25.00")
    assert len(invoice.items) == 1
    assert invoice.items[0].item_code.normalized_value == "000042"
    assert invoice.items[0].quantity.normalized_value == Decimal("2")
    assert invoice.items[0].unit_price.normalized_value == Decimal("12.5")
    assert invoice.items[0].amount.cell == "G18"
    assert invoice.items[0].amount.sheet == "Synthetic Invoice"
    assert any(candidate.field_path.startswith("items[") for candidate in outcome.result.candidates)
    restored = DocumentExtractionRepository().get_active(DOC)
    assert restored.result == outcome.result
    assert outcome.record == restored
    refreshed = SourceDocumentRepository().get_by_document_id(DOC)
    for name in ["file_hash", "stored_path", "detected_document_type", "classification_confidence",
                 "parser_used", "page_count", "sheet_count"]:
        assert getattr(refreshed, name) == getattr(source, name)
    assert path.read_bytes() == before


def test_real_native_pdf_preserves_page_and_explicitly_requires_item_review():
    path = _pdf()
    _source(path)
    outcome = _service().extract_document(DOC)
    assert outcome.persisted
    assert outcome.result.status == ExtractionStatus.NEEDS_REVIEW
    assert outcome.result.invoice.invoice_number.normalized_value == "000012"
    assert outcome.result.invoice.invoice_number.page == 1
    assert outcome.result.invoice.items == []
    assert any(issue.code == "PDF_ITEMS_UNSUPPORTED" for issue in outcome.result.issues)


def test_real_legacy_xls_uses_existing_parser_and_item_extractor():
    import xlwt
    path = _path("xls")
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("Synthetic")
    for row, values in [(0, ["COMMERCIAL INVOICE"]), (2, ["INVOICE NO", "000012"]),
                        (16, ["ITEM CODE", "QTY", "AMOUNT"]), (17, ["000042", 0, 0])]:
        for column, value in enumerate(values):
            sheet.write(row, column, value)
    workbook.save(str(path))
    before = path.read_bytes()
    _source(path)
    outcome = _service().extract_document(DOC)
    assert outcome.result.status == ExtractionStatus.EXTRACTED
    assert outcome.result.invoice.items[0].quantity.normalized_value == Decimal("0")
    assert outcome.result.invoice.items[0].amount.normalized_value == Decimal("0")
    assert outcome.result.invoice.items[0].item_code.normalized_value == "000042"
    assert path.read_bytes() == before


@pytest.mark.parametrize("kind", ["PACKING_LIST", "BILL_OF_LADING", "UNKNOWN", None])
def test_wrong_or_missing_stored_classification_refuses_without_semantic_or_write(kind):
    _source(_xlsx(), detected_document_type=kind)
    semantic = SemanticSpy()
    parsing = ReadSpy()
    outcome = _service(semantic_provider=semantic, parsing_service=parsing).extract_document(DOC)
    assert outcome.result.status == ExtractionStatus.NEEDS_REVIEW
    assert not outcome.persisted
    assert outcome.result.invoice.items == []
    assert semantic.requests == []
    assert parsing.calls == 0
    assert _history() == []


@pytest.mark.parametrize("confidence", [None, 0.1, float("inf")])
def test_missing_weak_or_invalid_stored_confidence_cannot_pass_gate(confidence):
    _source(_xlsx(), classification_confidence=confidence)
    outcome = _service().extract_document(DOC)
    assert outcome.result.status == ExtractionStatus.NEEDS_REVIEW
    assert not outcome.persisted
    assert _history() == []


@pytest.mark.parametrize("title", ["PACKING LIST", "BILL OF LADING", "COMMERCIAL INVOICE PACKING LIST", "Synthetic ordinary text"])
def test_actual_content_rechecks_classification_and_filename_does_not_override(title):
    source = _source(_xlsx(title=title, items=False))
    semantic = SemanticSpy()
    outcome = _service(semantic_provider=semantic).extract_document(DOC)
    assert outcome.result.status == ExtractionStatus.NEEDS_REVIEW
    assert not outcome.persisted
    assert semantic.requests == []
    assert SourceDocumentRepository().get_by_document_id(DOC) == source


@pytest.mark.parametrize("clue", ["formula", "hidden_row", "hidden_sheet"])
def test_unsafe_excel_classification_clues_cannot_authorize_semantic_or_extraction(clue):
    path = _xlsx(title="Synthetic ordinary text", items=False)
    workbook = openpyxl.load_workbook(path)
    if clue == "formula":
        workbook.active["A1"] = '="COMMERCIAL INVOICE"'
    elif clue == "hidden_row":
        workbook.active["A1"] = "COMMERCIAL INVOICE"
        workbook.active.row_dimensions[1].hidden = True
    else:
        hidden = workbook.create_sheet("Hidden clues")
        hidden["A1"] = "COMMERCIAL INVOICE"
        hidden.sheet_state = "hidden"
    workbook.save(path)
    workbook.close()
    source = _source(path)
    semantic = SemanticSpy()
    outcome = _service(semantic_provider=semantic).extract_document(DOC)
    assert not outcome.persisted
    assert outcome.result.status == ExtractionStatus.NEEDS_REVIEW
    assert semantic.requests == []
    assert SourceDocumentRepository().get_by_document_id(DOC) == source
    assert _history() == []


def test_hidden_header_value_cannot_be_published_or_semantically_recovered():
    path = _xlsx()
    workbook = openpyxl.load_workbook(path)
    workbook.active.row_dimensions[9].hidden = True
    workbook.save(path)
    workbook.close()
    _source(path)
    semantic = SemanticSpy(lambda request: SemanticInvoiceResponse(
        document_id=DOC, candidates=[ExtractionCandidate(
            field_path="total_amount", raw_value="25.00", typed_value=Decimal("25.00"),
            source_document_id=DOC, sheet="Synthetic Invoice", cell="B9",
        )],
    ))
    outcome = _service(semantic_provider=semantic).extract_document(DOC)
    assert outcome.persisted
    assert outcome.result.status == ExtractionStatus.NEEDS_REVIEW
    assert outcome.result.invoice.total_amount is None
    assert any(issue.code == "HIDDEN_INVOICE_CONTENT" for issue in outcome.result.issues)
    assert all(source.cell != "B9" for request in semantic.requests for source in request.sources)


@pytest.mark.parametrize("document_id", ["", " ", None, 42])
def test_malformed_document_identifier_returns_safe_typed_failure(document_id):
    outcome = _service().extract_document(document_id)
    assert outcome.result.status == ExtractionStatus.FAILED
    assert not outcome.persisted
    assert outcome.result.invoice.items == []


def test_missing_document_is_controlled_and_does_not_create_history():
    outcome = _service().extract_document("DOC-MISSING")
    assert outcome.result.status == ExtractionStatus.FAILED
    assert not outcome.persisted
    assert DocumentExtractionRepository().list_history("DOC-MISSING") == []


@pytest.mark.parametrize("initial_status", [ProcessingStatus.CLASSIFIED, ProcessingStatus.NEEDS_REVIEW])
def test_parsing_does_not_downgrade_classified_reviewed_or_extracted_source(initial_status):
    source = _source(_xlsx(), processing_status=initial_status)
    parsing = ReadSpy()
    service = _service(parsing_service=parsing)
    first = service.extract_document(DOC)
    assert parsing.observations == [source, source]
    extracted = SourceDocumentRepository().get_by_document_id(DOC)
    second = service.extract_document(DOC, expected_active_extraction_id=first.record.extraction_id)
    assert second.persisted
    assert parsing.observations[2:] == [extracted, extracted]
    assert extracted.processing_status == ProcessingStatus.EXTRACTED
    assert len(_history()) == 2


def test_source_hash_mismatch_refuses_without_changing_metadata_or_active():
    path = _xlsx()
    source = _source(path)
    path.write_bytes(path.read_bytes() + b"synthetic external change")
    outcome = _service().extract_document(DOC)
    assert not outcome.persisted
    assert outcome.result.status in {ExtractionStatus.FAILED, ExtractionStatus.NEEDS_REVIEW}
    assert SourceDocumentRepository().get_by_document_id(DOC) == source
    assert _history() == []


def test_source_path_containment_is_enforced_before_parsing(tmp_path):
    outside = tmp_path / "outside.xlsx"
    outside.write_bytes(_xlsx().read_bytes())
    source = _source(outside)
    outcome = _service().extract_document(DOC)
    assert outcome.result.status == ExtractionStatus.FAILED
    assert not outcome.persisted
    assert SourceDocumentRepository().get_by_document_id(DOC) == source


@pytest.mark.parametrize("change", ["file", "status", "classification"])
def test_source_changes_during_processing_are_not_persisted(change):
    path = _xlsx()
    _source(path)
    def mutate(parsed):
        if change == "file":
            path.write_bytes(path.read_bytes() + b"synthetic changed during processing")
        else:
            SourceDocumentRepository().update_classification_metadata(
                DOC, "PACKING_LIST" if change == "classification" else "COMMERCIAL_INVOICE",
                0.9, ProcessingStatus.NEEDS_REVIEW if change == "status" else ProcessingStatus.CLASSIFIED,
            )
    outcome = _service(parsing_service=ReadSpy(mutate)).extract_document(DOC)
    assert not outcome.persisted
    assert _history() == []


def test_scanned_pdf_without_vision_is_reviewed_and_never_sent_to_semantic():
    _source(_pdf(blank=True))
    semantic = SemanticSpy()
    outcome = _service(semantic_provider=semantic).extract_document(DOC)
    assert outcome.result.status == ExtractionStatus.NEEDS_REVIEW
    assert outcome.persisted
    assert outcome.result.invoice.invoice_number is None
    assert semantic.requests == []


@pytest.mark.parametrize("unsafe", ["document", "page", "unresolved", "unverified", "low_confidence", "mixed"])
def test_unsafe_visual_resolution_never_reaches_semantic_or_trusted_extraction(unsafe):
    _source(_pdf(blank=True, pages=2 if unsafe == "mixed" else 1))
    page = ResolvedPdfPage(
        page=1, text="COMMERCIAL INVOICE\nINVOICE NO: 000012", confidence=0.99,
        source=TextResolutionSource.CROSS_VERIFIED_VISUAL,
        verification_level=VerificationLevel.CROSS_VERIFIED, requires_review=False,
    )
    if unsafe == "page":
        page.requires_review = True
    if unsafe == "unverified":
        page.verification_level = VerificationLevel.UNVERIFIED
    if unsafe == "low_confidence":
        page.confidence = 0.01
    pages = [page]
    if unsafe == "mixed":
        pages.append(ResolvedPdfPage(page=2, text="", source=TextResolutionSource.UNRESOLVED,
                                     verification_level=VerificationLevel.UNVERIFIED, requires_review=True))
    resolved = ResolvedPdfDocument(
        document_id=DOC, page_count=len(pages), pages=pages, visual_pages=[p.page for p in pages],
        unresolved_pages=[1] if unsafe == "unresolved" else [], requires_review=unsafe == "document",
    )
    class Vision:
        calls = 0
        def resolve_document(self, document_id, parsed):
            self.calls += 1
            return resolved
    vision = Vision()
    semantic = SemanticSpy()
    outcome = _service(vision_service=vision, semantic_provider=semantic).extract_document(DOC)
    assert outcome.result.status == ExtractionStatus.NEEDS_REVIEW
    assert outcome.result.invoice.invoice_number is None
    assert semantic.requests == []
    assert vision.calls == 1


def test_safe_resolved_pdf_uses_verified_page_once_and_retains_pdf_item_limitation():
    _source(_pdf(blank=True))
    resolved = ResolvedPdfDocument(
        document_id=DOC, page_count=1, requires_review=False, visual_pages=[1],
        pages=[ResolvedPdfPage(
            page=1, text="COMMERCIAL INVOICE\nINVOICE NO: 000012", confidence=0.99,
            source=TextResolutionSource.CROSS_VERIFIED_VISUAL,
            verification_level=VerificationLevel.CROSS_VERIFIED, requires_review=False,
        )],
    )
    class Vision:
        calls = 0
        def resolve_document(self, document_id, parsed):
            self.calls += 1
            return resolved
    vision = Vision()
    semantic = SemanticSpy()
    outcome = _service(vision_service=vision, semantic_provider=semantic).extract_document(DOC)
    assert outcome.persisted
    assert outcome.result.invoice.invoice_number.normalized_value == "000012"
    assert outcome.result.invoice.invoice_number.page == 1
    assert outcome.result.status == ExtractionStatus.NEEDS_REVIEW
    assert any(issue.code == "PDF_ITEMS_UNSUPPORTED" for issue in outcome.result.issues)
    assert vision.calls == 1


def test_vision_failure_is_review_and_does_not_fabricate_business_facts():
    _source(_pdf(blank=True))
    class Vision:
        def resolve_document(self, document_id, parsed):
            raise RuntimeError("SYNTHETIC_PRIVATE_PROVIDER_FAILURE")
    semantic = SemanticSpy()
    outcome = _service(vision_service=Vision(), semantic_provider=semantic).extract_document(DOC)
    assert outcome.result.status == ExtractionStatus.NEEDS_REVIEW
    assert outcome.result.invoice.invoice_number is None
    assert semantic.requests == []
    assert "SYNTHETIC_PRIVATE_PROVIDER_FAILURE" not in str(outcome.result.issues)


def test_native_pdf_does_not_invoke_unnecessary_vision():
    _source(_pdf())
    class Vision:
        def resolve_document(self, *args):
            raise AssertionError("Native safe PDF should bypass vision")
    outcome = _service(vision_service=Vision()).extract_document(DOC)
    assert outcome.result.invoice.invoice_number.normalized_value == "000012"
    assert not any("VISION" in issue.code for issue in outcome.result.issues)


@pytest.mark.parametrize("incomplete,items,extra_table", [(True, True, False), (False, False, False), (False, True, True)])
def test_partial_missing_or_competing_item_tables_remain_review(incomplete, items, extra_table):
    _source(_xlsx(incomplete=incomplete, items=items, extra_table=extra_table))
    outcome = _service().extract_document(DOC)
    assert outcome.result.status == ExtractionStatus.NEEDS_REVIEW
    assert outcome.persisted
    assert any(issue.is_review_required for issue in outcome.result.issues)
    if incomplete:
        assert len(outcome.result.invoice.items) == 2
        assert outcome.result.invoice.items[1].quantity is None
    else:
        assert outcome.result.invoice.items == []


def test_formula_uncertainty_is_sticky_despite_semantic_response():
    _source(_xlsx(formula=True))
    semantic = SemanticSpy(lambda request: SemanticInvoiceResponse(document_id=DOC, candidates=[ExtractionCandidate(
        field_path="total_amount", raw_value="=G18", typed_value=Decimal("25.00"),
        source_document_id=DOC, sheet="Synthetic Invoice", cell="B9",
    )]))
    outcome = _service(semantic_provider=semantic).extract_document(DOC)
    assert outcome.result.invoice.total_amount is None
    assert outcome.result.status == ExtractionStatus.NEEDS_REVIEW
    assert any(issue.code == "UNRESOLVED_FORMULA_FIELD" for issue in outcome.result.issues)
    assert any(candidate.raw_value == "=G18" for candidate in outcome.result.candidates)


def test_deterministic_conflict_cannot_be_resolved_by_semantic_single_value():
    _source(_xlsx(conflict=True))
    semantic = SemanticSpy(lambda request: SemanticInvoiceResponse(document_id=DOC, candidates=[ExtractionCandidate(
        field_path="invoice_number", raw_value="000012", typed_value="000012",
        source_document_id=DOC, sheet="Synthetic Invoice", cell="B3",
    )]))
    outcome = _service(semantic_provider=semantic).extract_document(DOC)
    assert outcome.result.invoice.invoice_number is None
    assert outcome.result.status == ExtractionStatus.NEEDS_REVIEW
    assert any(issue.code == "CONFLICTING_FIELD_VALUES" for issue in outcome.result.issues)


def test_semantic_provider_failure_is_review_and_sanitized(caplog):
    _source(_xlsx())
    def fail(request):
        raise RuntimeError("SYNTHETIC_PRIVATE_SEMANTIC_PAYLOAD")
    outcome = _service(semantic_provider=SemanticSpy(fail)).extract_document(DOC)
    assert outcome.result.status == ExtractionStatus.NEEDS_REVIEW
    assert outcome.persisted
    assert outcome.result.invoice.invoice_number.normalized_value == "000012"
    assert "SYNTHETIC_PRIVATE_SEMANTIC_PAYLOAD" not in str(outcome.result.issues) + caplog.text


def test_wrong_semantic_document_is_review_without_corrupting_deterministic_value():
    _source(_xlsx())
    provider = SemanticSpy(lambda request: SemanticInvoiceResponse(document_id="DOC-OTHER", candidates=[]))
    outcome = _service(semantic_provider=provider).extract_document(DOC)
    assert outcome.result.status == ExtractionStatus.NEEDS_REVIEW
    assert outcome.result.invoice.invoice_number.normalized_value == "000012"


def test_valid_grounded_semantic_fill_of_missing_field_preserves_other_values():
    _source(_xlsx(freight=True))
    class MissingFreight:
        def extract(self, document):
            result = DeterministicInvoiceHeaderExtractor().extract(document)
            result.invoice.freight = None
            result.candidates = [candidate for candidate in result.candidates if candidate.field_path != "freight"]
            return result
    provider = SemanticSpy(lambda request: SemanticInvoiceResponse(document_id=DOC, candidates=[ExtractionCandidate(
        field_path="freight", raw_value="2.50", typed_value=Decimal("2.50"),
        source_document_id=DOC, sheet="Synthetic Invoice", cell="B12",
    )]))
    outcome = _service(header_extractor=MissingFreight(), semantic_provider=provider).extract_document(DOC)
    assert outcome.result.invoice.freight.normalized_value == Decimal("2.50")
    assert outcome.result.invoice.freight.cell == "B12"
    assert outcome.result.invoice.invoice_number.normalized_value == "000012"
    assert outcome.persisted
    assert outcome.result.status == ExtractionStatus.EXTRACTED


def test_existing_active_requires_explicit_reextraction_and_preserves_history():
    _source(_xlsx())
    service = _service()
    first = service.extract_document(DOC)
    rejected = service.extract_document(DOC)
    assert not rejected.persisted
    assert _history() == [first.record]
    second = service.extract_document(DOC, expected_active_extraction_id=first.record.extraction_id)
    assert second.persisted
    assert [record.is_active for record in _history()] == [False, True]
    assert second.record.supersedes_extraction_id == first.record.extraction_id


@pytest.mark.parametrize("precondition", ["active", "hash", "status"])
def test_stale_reextraction_does_not_change_prior_active_or_source(precondition):
    _source(_xlsx())
    service = _service()
    first = service.extract_document(DOC)
    source_before = SourceDocumentRepository().get_by_document_id(DOC)
    arguments = {"expected_active_extraction_id": first.record.extraction_id}
    arguments.update({
        "active": {"expected_active_extraction_id": "missing-active"},
        "hash": {"expected_source_file_hash": "f" * 64},
        "status": {"expected_processing_status": ProcessingStatus.CLASSIFIED},
    }[precondition])
    rejected = service.extract_document(DOC, **arguments)
    assert not rejected.persisted
    assert DocumentExtractionRepository().get_active(DOC) == first.record
    assert _history() == [first.record]
    assert SourceDocumentRepository().get_by_document_id(DOC) == source_before


def test_failed_later_attempt_preserves_active_and_latest_status_is_distinct():
    _source(_xlsx())
    first = _service().extract_document(DOC)
    class FailedParser:
        def parse_document_readonly(self, document_id):
            raise ParserFailedError()
    failed = _service(parsing_service=FailedParser()).extract_document(
        DOC, expected_active_extraction_id=first.record.extraction_id,
    )
    assert failed.result.status == ExtractionStatus.FAILED
    assert failed.persisted
    assert failed.record.is_active is False
    assert DocumentExtractionRepository().get_active(DOC) == first.record
    latest_source = SourceDocumentRepository().get_by_document_id(DOC)
    assert latest_source.processing_status == ProcessingStatus.FAILED
    assert latest_source.extraction_status == "FAILED"
    assert len(_history()) == 2


@pytest.mark.parametrize("replacement", [False, True])
def test_classification_change_at_repository_boundary_is_rejected_atomically(replacement):
    _source(_xlsx())
    first = _service().extract_document(DOC) if replacement else None

    class RacingRepository(DocumentExtractionRepository):
        def change_classification(self):
            source = SourceDocumentRepository().get_by_document_id(DOC)
            # Hash and processing status stay unchanged: original 006D guards
            # alone cannot notice this independently material metadata change.
            SourceDocumentRepository().update_classification_metadata(
                DOC, "PACKING_LIST", 0.99, source.processing_status,
            )

        def save_attempt(self, attempt):
            self.change_classification()
            return super().save_attempt(attempt)

        def replace_active(self, attempt):
            self.change_classification()
            return super().replace_active(attempt)

    arguments = {"expected_active_extraction_id": first.record.extraction_id} if first else {}
    outcome = _service(extraction_repository=RacingRepository()).extract_document(DOC, **arguments)
    assert outcome.result.status == ExtractionStatus.FAILED
    assert not outcome.persisted
    assert len(_history()) == (1 if first else 0)
    assert DocumentExtractionRepository().get_active(DOC) == (first.record if first else None)
    assert SourceDocumentRepository().get_by_document_id(DOC).detected_document_type == "PACKING_LIST"


def test_persistence_failure_rolls_back_and_returns_safe_unpersisted_failure():
    _source(_xlsx())
    first = _service().extract_document(DOC)
    source_before = SourceDocumentRepository().get_by_document_id(DOC)
    with closing(get_connection()) as connection, connection:
        connection.execute("""CREATE TRIGGER synthetic_abort BEFORE UPDATE OF extraction_status
            ON source_documents BEGIN SELECT RAISE(ABORT, 'SYNTHETIC_PRIVATE_DB_ERROR'); END""")
    failed = _service().extract_document(DOC, expected_active_extraction_id=first.record.extraction_id)
    assert failed.result.status == ExtractionStatus.FAILED
    assert not failed.persisted
    assert "SYNTHETIC_PRIVATE_DB_ERROR" not in str(failed.result.issues)
    assert DocumentExtractionRepository().get_active(DOC) == first.record
    assert _history() == [first.record]
    assert SourceDocumentRepository().get_by_document_id(DOC) == source_before


def test_source_content_and_sensitive_paths_are_not_logged(caplog):
    path = _xlsx()
    _source(path)
    outcome = _service().extract_document(DOC)
    assert outcome.persisted
    for sensitive in [str(path), "Synthetic Seller", "Synthetic Buyer", "Synthetic Motor", "000012"]:
        assert sensitive not in caplog.text


def test_production_factory_works_without_semantic_or_network_calls(monkeypatch):
    import socket
    def no_network(*args, **kwargs):
        raise AssertionError("No network call permitted")
    monkeypatch.setattr(socket.socket, "connect", no_network)
    monkeypatch.setattr(settings.vision, "enabled", False)
    _source(_xlsx())
    service = build_invoice_extraction_service(SourceDocumentRepository())
    assert isinstance(service, InvoiceExtractionService)
    assert service.extract_document(DOC).result.status == ExtractionStatus.EXTRACTED
