"""Actual parser integration and final source-boundary regressions for 007C."""

from decimal import Decimal
from hashlib import sha256

import pytest
import xlwt

from customs_ai.extraction.enums import ExtractionStatus
from customs_ai.extraction.packing_items import DeterministicPackingItemExtractor, extract_packing_items
from customs_ai.parsers.excel import XlsParser
from customs_ai.parsers.models import ParsedCell, ParsedPdfDocument, ParsedPdfPage, ParsedSheet, ParsedWorkbook


def _workbook(rows, *, hidden_rows=()):
    from openpyxl.utils import get_column_letter
    return ParsedWorkbook(
        document_id="PACK-INTEGRATION", file_type="xlsx", parser_used="openpyxl",
        workbook_name="synthetic.xlsx", sheet_count=1,
        sheets=[ParsedSheet(name="Synthetic", index=0, state="visible", hidden_rows=list(hidden_rows),
                            cells=[ParsedCell(coordinate=f"{get_column_letter(col)}{row}", row=row,
                                              column=col, raw_value=value, data_type="s")
                                   for row, values in enumerate(rows, start=1)
                                   for col, value in enumerate(values, start=1) if value is not None])],
    )


def test_actual_xls_preserves_all_item_fields_source_hash_and_coordinates(tmp_path):
    path = tmp_path / "synthetic-packing.xls"
    book = xlwt.Workbook()
    sheet = book.add_sheet("Synthetic XLS")
    labels = ["ITEM CODE", "DESCRIPTION", "QTY", "UNIT", "PACKAGE NO.", "GROSS WEIGHT", "NET WEIGHT"]
    values = ["000012", "Synthetic Pump", "0.00", "PCS", "000045", "12345678901234567890.1200", "-2.50"]
    for col, (label, value) in enumerate(zip(labels, values), start=4):
        sheet.write(16, col, label)
        sheet.write(17, col, value)
    book.save(str(path))
    before = sha256(path.read_bytes()).hexdigest()
    parsed = XlsParser().parse("PACK-XLS", path)
    snapshot = parsed.model_dump(mode="python")
    result = extract_packing_items(parsed)
    assert result.status == ExtractionStatus.EXTRACTED and len(result.items) == 1
    item = result.items[0]
    for field, raw, coordinate in zip(
            [item.item_code, item.description, item.quantity, item.unit, item.package_number,
             item.gross_weight, item.net_weight], values, ["E18", "F18", "G18", "H18", "I18", "J18", "K18"]):
        assert field.raw_value == raw
        assert field.source_document_id == "PACK-XLS" and field.sheet == "Synthetic XLS"
        assert field.cell == coordinate and field.page is None and field.extraction_method
    assert item.item_code.normalized_value == "000012" and item.package_number.normalized_value == "000045"
    assert item.quantity.normalized_value.as_tuple() == Decimal("0.00").as_tuple()
    assert item.gross_weight.normalized_value.as_tuple() == Decimal(values[5]).as_tuple()
    assert item.net_weight.normalized_value == Decimal("-2.50")
    assert sha256(path.read_bytes()).hexdigest() == before
    assert parsed.model_dump(mode="python") == snapshot


@pytest.mark.parametrize("document", [None, object(), ParsedPdfDocument(
    document_id="SYNTHETIC-PDF", page_count=1, has_text_layer=True, needs_ocr=False,
    pages=[ParsedPdfPage(page=1, text="DESCRIPTION | QTY\nSynthetic Pump | 2")],
)])
def test_non_workbook_input_never_attempts_pdf_item_extraction(document):
    result = extract_packing_items(document)
    assert result.status == ExtractionStatus.FAILED and result.items == []
    assert result.candidates == [] and result.issues


@pytest.mark.parametrize("marker", ["NOTES", "REMARKS", "GHI CHÚ", "备注"])
def test_bare_note_heading_does_not_certify_a_contiguous_example_table(marker):
    result = extract_packing_items(_workbook([[marker], ["DESCRIPTION", "QTY"], ["Example gear", "2"]]))
    assert result.status == ExtractionStatus.NEEDS_REVIEW and result.items == []
    assert any(candidate.raw_value == "Example gear" for candidate in result.candidates)


def test_hidden_empty_gap_in_table_does_not_claim_clean_complete_extraction():
    result = extract_packing_items(_workbook([
        ["DESCRIPTION", "QTY"], ["Synthetic Pump", "2"], [], ["Synthetic Gear", "3"],
    ], hidden_rows=[3]))
    assert result.status == ExtractionStatus.NEEDS_REVIEW
    assert len(result.items) == 1 and result.items[0].description.normalized_value == "Synthetic Pump"
    assert any(candidate.raw_value == "Synthetic Gear" for candidate in result.candidates)


def test_vertical_merge_preserves_real_fields_on_all_affected_rows_without_distribution():
    document = _workbook([
        ["ITEM CODE", "DESCRIPTION", "QTY"], ["0001", "Safe gear", "1"],
        ["0002", "Shared merged description", "2"], ["0003", None, "3"],
    ])
    document.sheets[0].merged_ranges = ["B3:B4"]
    result = extract_packing_items(document)
    assert result.status == ExtractionStatus.NEEDS_REVIEW and len(result.items) == 1
    assert {candidate.raw_value for candidate in result.candidates} >= {"0002", "0003", "2", "3"}
    assert all(candidate.cell != "B4" for candidate in result.candidates)


@pytest.mark.parametrize("marker", ["TOTAL : 99", "TOTAL PACKAGES : 99", "NOTES : synthetic note"])
def test_explicit_footer_colon_with_whitespace_never_creates_another_product(marker):
    result = extract_packing_items(_workbook([
        ["DESCRIPTION", "QTY"], ["Synthetic Gear", "2"], [marker, "99"],
    ]))
    assert len(result.items) == 1
    assert result.items[0].description.normalized_value == "Synthetic Gear"
    assert any(candidate.raw_value == marker for candidate in result.candidates)
