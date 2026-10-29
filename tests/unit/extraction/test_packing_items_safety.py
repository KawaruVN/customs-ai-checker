"""TASK-007C adversarial source safety, geometry and actual-XLSX probes."""

from datetime import date, datetime, time
from decimal import Decimal
from hashlib import sha256

from openpyxl import Workbook
from openpyxl.utils.cell import get_column_letter
import pytest

from customs_ai.domain.enums import ValueOrigin
from customs_ai.extraction.enums import ExtractionStatus
from customs_ai.extraction.packing_config import PackingExtractionConfig, load_packing_extraction_config
from customs_ai.extraction.packing_items import DeterministicPackingItemExtractor
from customs_ai.parsers.excel import XlsxParser
from customs_ai.parsers.models import ParsedCell, ParsedSheet, ParsedWorkbook


def _cell(row, col, value, **metadata):
    return ParsedCell(
        coordinate=f"{get_column_letter(col)}{row}", row=row, column=col,
        raw_value=value, data_type=metadata.pop("data_type", "s"), **metadata,
    )


def _sheet(rows, *, row=1, col=1, name="Synthetic Packing", **metadata):
    return ParsedSheet(
        name=name, index=metadata.pop("index", 0), state=metadata.pop("state", "visible"),
        cells=[_cell(row + i, col + j, value) for i, values in enumerate(rows)
               for j, value in enumerate(values) if value is not None], **metadata,
    )


def _book(*sheets):
    return ParsedWorkbook(
        document_id="PACK-ITEM-SAFETY", file_type="xlsx", parser_used="openpyxl",
        workbook_name="synthetic.xlsx", sheet_count=len(sheets),
        sheets=[sheet.model_copy(update={"index": i}) for i, sheet in enumerate(sheets)],
    )


def _extract(*sheets):
    return DeterministicPackingItemExtractor().extract(_book(*sheets))


def _review(result):
    assert result.status == ExtractionStatus.NEEDS_REVIEW
    assert result.issues and any(issue.is_review_required for issue in result.issues)


def _normal():
    return _sheet([
        ["ITEM CODE", "DESCRIPTION", "QUANTITY", "UNIT", "PACKAGE NO.", "GROSS WEIGHT", "NET WEIGHT"],
        ["000012", "Synthetic Gear", "2", "PCS", "000045", "12.50", "11.50"],
        ["000013", "Synthetic Cable", "3", "PCS", "000046", "13.50", "12.50"],
    ])


@pytest.mark.parametrize("raw", [
    "1,200", "1.200", "KG 10", "10 KG", "NaN", "Infinity", True,
    date(2026, 10, 29), datetime(2026, 10, 29, 12, 30), time(12, 30),
    float("nan"), float("inf"), float("-inf"),
])
@pytest.mark.parametrize("path,col", [("quantity", 3), ("gross_weight", 6), ("net_weight", 7)])
def test_unsafe_numeric_cell_is_retained_only_as_candidate(raw, path, col):
    sheet = _normal()
    target = next(cell for cell in sheet.cells if (cell.row, cell.column) == (3, col))
    target.raw_value = raw
    result = _extract(sheet)
    _review(result)
    assert result.items and result.items[0].item_code.normalized_value == "000012"
    affected = [item for item in result.items if item.item_code is not None and item.item_code.normalized_value == "000013"]
    if affected:
        assert getattr(affected[0], path) is None
    assert any(candidate.cell == target.coordinate and candidate.raw_value == str(raw) for candidate in result.candidates)


@pytest.mark.parametrize("path,col", [("quantity", 3), ("gross_weight", 6), ("net_weight", 7)])
@pytest.mark.parametrize("metadata,raw", [
    ({"is_formula": True}, "=2+2"),
    ({"data_type": "f"}, "4"),
    ({"is_formula": True, "displayed_value": 4}, "4"),
    ({"displayed_value": 4}, " =2+2"),
])
def test_formula_quantities_and_weights_never_trust_expression_or_cached_value(path, col, metadata, raw):
    sheet = _normal()
    old = next(cell for cell in sheet.cells if (cell.row, cell.column) == (3, col))
    sheet.cells[sheet.cells.index(old)] = _cell(3, col, raw, **metadata)
    result = _extract(sheet)
    _review(result)
    for item in result.items:
        value = getattr(item, path)
        assert value is None or value.cell != old.coordinate
    assert any(candidate.cell == old.coordinate and candidate.raw_value == raw for candidate in result.candidates)


@pytest.mark.parametrize("col,path", [(1, "item_code"), (2, "description"), (4, "unit"), (5, "package_number")])
def test_formula_text_fields_remain_unresolved_even_with_displayed_cache(col, path):
    sheet = _normal()
    old = next(cell for cell in sheet.cells if (cell.row, cell.column) == (3, col))
    sheet.cells[sheet.cells.index(old)] = _cell(3, col, '="FAKE-CACHE"', is_formula=True, displayed_value="FAKE-CACHE")
    result = _extract(sheet)
    _review(result)
    assert all(getattr(item, path) is None or getattr(item, path).cell != old.coordinate for item in result.items)
    assert any(candidate.cell == old.coordinate and candidate.raw_value == '="FAKE-CACHE"' for candidate in result.candidates)


@pytest.mark.parametrize("metadata,raw", [
    ({"is_formula": True}, "DESCRIPTION"),
    ({"data_type": "f"}, "DESCRIPTION"),
    ({"displayed_value": "DESCRIPTION"}, '= "DESCRIPTION"'),
])
def test_formula_header_label_never_establishes_product_table(metadata, raw):
    sheet = _sheet([["DESCRIPTION", "QUANTITY"], ["Synthetic Gear", "2"]])
    sheet.cells[0] = _cell(1, 1, raw, **metadata)
    result = _extract(sheet)
    _review(result)
    assert result.items == []


@pytest.mark.parametrize("path,col", [("item_code", 1), ("package_number", 5)])
@pytest.mark.parametrize("raw", [12, 12.0, True])
def test_numeric_identifier_does_not_invent_display_zeroes(path, col, raw):
    sheet = _normal()
    old = next(cell for cell in sheet.cells if (cell.row, cell.column) == (3, col))
    sheet.cells[sheet.cells.index(old)] = _cell(3, col, raw, number_format="000000")
    result = _extract(sheet)
    _review(result)
    assert all(getattr(item, path) is None or getattr(item, path).cell != old.coordinate for item in result.items)
    assert any(candidate.cell == old.coordinate and candidate.raw_value == str(raw) for candidate in result.candidates)


@pytest.mark.parametrize("hidden", ["A", "B", "C", "D", "E", "F", "G"])
def test_mapped_hidden_column_blocks_affected_table(hidden):
    result = _extract(_normal().model_copy(update={"hidden_columns": [hidden]}))
    _review(result)
    assert result.items == []


@pytest.mark.parametrize("state", ["hidden", "veryHidden"])
def test_hidden_worksheet_never_contributes_items_or_outvotes_visible_table(state):
    hidden = _normal().model_copy(update={"name": "Hidden", "state": state})
    visible = _sheet([["DESCRIPTION", "QUANTITY"], ["Visible Widget", "8"]], name="Visible")
    result = _extract(hidden, visible)
    _review(result)
    assert len(result.items) == 1
    assert result.items[0].description.normalized_value == "Visible Widget"
    assert result.items[0].description.sheet == "Visible"


def test_hidden_mapped_column_does_not_outvote_an_independent_safe_table():
    unsafe = _sheet(
        [["DESCRIPTION", "QUANTITY", "PACKAGE NO."], ["Unsafe Gear", "2", "000045"]],
        name="Unsafe", hidden_columns=["C"],
    )
    safe = _sheet([["DESCRIPTION", "QUANTITY"], ["Safe Gear", "3"]], name="Safe")
    result = _extract(unsafe, safe)
    _review(result)
    assert len(result.items) == 1
    assert result.items[0].description.normalized_value == "Safe Gear"
    assert result.items[0].description.sheet == "Safe"


@pytest.mark.parametrize("footer", [
    "TOTAL GROSS WEIGHT", "TOTAL NET WEIGHT", "TOTAL PACKAGES",
    "TỔNG TRỌNG LƯỢNG CẢ BÌ", "总毛重",
])
def test_explicit_document_total_label_is_not_an_additional_product(footer):
    result = _extract(_sheet([
        ["DESCRIPTION", "QUANTITY", "GROSS WEIGHT"],
        ["Safe Gear", "2", "12.50"],
        [footer, "3", "99.50"],
    ]))
    assert len(result.items) == 1
    assert result.items[0].description.normalized_value == "Safe Gear"
    assert any(candidate.raw_value == footer for candidate in result.candidates)


def test_hidden_header_cannot_supply_table_mapping():
    result = _extract(_normal().model_copy(update={"hidden_rows": [1]}))
    _review(result)
    assert result.items == []


def test_hidden_business_row_is_omitted_with_review_and_source_candidates():
    result = _extract(_normal().model_copy(update={"hidden_rows": [3]}))
    _review(result)
    assert len(result.items) == 1 and result.items[0].item_code.normalized_value == "000012"
    assert all(field is None or field.cell not in {"A3", "B3", "C3", "D3", "E3", "F3", "G3"}
               for item in result.items for _, field in item)
    assert any(candidate.cell == "A3" and candidate.raw_value == "000013" for candidate in result.candidates)


@pytest.mark.parametrize("merge", ["A1:B1", "B1:C1", "C1:D1", "A1:A2"])
def test_merged_header_mapping_is_not_guessed(merge):
    # Use only real anchor values; metadata can still make unique columns unsafe.
    sheet = _sheet([["DESCRIPTION", None, "QUANTITY"], ["Synthetic Gear", None, "2"]], merged_ranges=[merge])
    result = _extract(sheet)
    _review(result)
    assert result.items == []


def test_horizontal_description_merge_preserves_actual_anchor_without_ghosts():
    result = _extract(_sheet(
        [["DESCRIPTION", None, "QUANTITY"], ["Synthetic Gear", None, "2"]],
        merged_ranges=["A2:B2"],
    ))
    assert result.status == ExtractionStatus.EXTRACTED
    assert len(result.items) == 1
    assert result.items[0].description.cell == "A2"
    assert result.items[0].quantity.cell == "C2"
    assert not any(candidate.cell == "B2" for candidate in result.candidates)


@pytest.mark.parametrize("merge", ["A3:A4", "B3:B4", "A3:B3"])
def test_vertical_or_cross_field_merge_cannot_distribute_values(merge):
    rows = [["DESCRIPTION", "QUANTITY"], ["Safe Gear", "2"], ["Merged Part", None], [None, "3"]]
    if merge == "B3:B4":
        rows[2][1] = "3"
        rows[3][1] = None
    sheet = _sheet(
        rows,
        merged_ranges=[merge],
    )
    result = _extract(sheet)
    _review(result)
    assert result.items and result.items[0].description.normalized_value == "Safe Gear"
    assert all(item.description is None or item.description.raw_value != "Merged Part" for item in result.items)
    assert not any(item.quantity is not None and item.quantity.cell == "B4" for item in result.items)


@pytest.mark.parametrize("ranges", [
    ["Z1:A1"], ["A0:B1"], ["not-a-range"], ["A1:B2", "B2:C3"], ["A1:A2", "A1:A2"],
])
def test_inconsistent_merge_metadata_abstains_without_crash(ranges):
    result = _extract(_normal().model_copy(update={"merged_ranges": ranges}))
    _review(result)
    assert result.items == []


def test_populated_merged_interior_is_not_allowed_to_fake_actual_evidence():
    sheet = _normal().model_copy(update={"merged_ranges": ["B2:C2"]})
    result = _extract(sheet)
    _review(result)
    assert result.items == []


@pytest.mark.parametrize("patch", [{"coordinate": "Z99"}, {"row": 0}, {"column": 0}, {"row": 9}, {"column": 9}])
def test_cell_coordinate_metadata_must_match_actual_grid_identity(patch):
    sheet = _normal()
    sheet.cells[0] = sheet.cells[0].model_copy(update=patch)
    result = _extract(sheet)
    _review(result)
    assert result.items == []


def test_duplicate_grid_coordinate_cannot_overwrite_source_evidence():
    sheet = _normal()
    sheet.cells.append(_cell(2, 1, "CONFLICTING-SECRET"))
    result = _extract(sheet)
    _review(result)
    assert result.items == []
    assert all("CONFLICTING-SECRET" not in issue.message for issue in result.issues)


@pytest.mark.parametrize("patch", [{"sheet_count": 8}, {"document_id": ""}, {"document_id": "  "}])
def test_invalid_workbook_identity_or_count_is_not_trusted(patch):
    book = _book(_normal()).model_copy(update=patch)
    result = DeterministicPackingItemExtractor().extract(book)
    assert result.status != ExtractionStatus.EXTRACTED and result.items == []
    assert result.issues


def test_source_snapshot_remains_unchanged_and_candidate_paths_use_physical_rows():
    book = _book(_sheet([["DESCRIPTION", "QUANTITY"], ["  MiXeD Synthetic Gear  ", "0.00"]], row=17, col=5))
    before = book.model_dump(mode="python")
    result = DeterministicPackingItemExtractor().extract(book)
    assert result.status == ExtractionStatus.EXTRACTED
    assert book.model_dump(mode="python") == before
    assert result.items[0].description.raw_value == "  MiXeD Synthetic Gear  "
    assert result.items[0].description.cell == "E18"
    assert result.items[0].quantity.normalized_value.as_tuple() == Decimal("0.00").as_tuple()
    assert result.items[0].description.origin == ValueOrigin.EXTRACTED
    assert result.items[0].description.source_document_id == "PACK-ITEM-SAFETY"
    assert result.items[0].description.extraction_method and result.items[0].description.confidence
    assert {candidate.field_path for candidate in result.candidates if not candidate.field_path.startswith("headers.")} == {
        "tables[0:17].rows[18].columns[5].description", "tables[0:17].rows[18].columns[6].quantity",
    }


def test_actual_xlsx_group_hidden_columns_remain_unsafe_and_source_bytes_unchanged(tmp_path):
    path = tmp_path / "synthetic-packing-hidden-group.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Synthetic"
    for row, values in enumerate([
        ["ITEM CODE", "DESCRIPTION", "QUANTITY", "PACKAGE NO."],
        ["000012", "Synthetic Gear", 2, "000045"],
    ], start=1):
        for col, value in enumerate(values, start=1):
            sheet.cell(row, col, value)
    sheet.column_dimensions.group("B", "D", hidden=True)
    book.save(path)
    book.close()
    digest = sha256(path.read_bytes()).hexdigest()
    parsed = XlsxParser().parse("PACK-XLSX", path)
    assert parsed.sheets[0].hidden_columns == ["B", "C", "D"]
    result = DeterministicPackingItemExtractor().extract(parsed)
    _review(result)
    assert result.items == []
    assert sha256(path.read_bytes()).hexdigest() == digest


def test_actual_xlsx_formula_keeps_expression_instead_of_cached_display(tmp_path):
    path = tmp_path / "synthetic-packing-formula.xlsx"
    book = Workbook()
    sheet = book.active
    for row in [["DESCRIPTION", "QUANTITY"], ["Safe Gear", 2], ["Formula Gear", "=1+1"]]:
        sheet.append(row)
    book.save(path)
    book.close()
    digest = sha256(path.read_bytes()).hexdigest()
    parsed = XlsxParser().parse("PACK-XLSX", path)
    result = DeterministicPackingItemExtractor().extract(parsed)
    _review(result)
    assert result.items[0].description.normalized_value == "Safe Gear"
    assert any(candidate.raw_value == "=1+1" and candidate.cell == "B3" for candidate in result.candidates)
    assert all(item.quantity is None or item.quantity.cell != "B3" for item in result.items)
    assert sha256(path.read_bytes()).hexdigest() == digest


def test_source_identifiers_and_description_never_leak_into_issue_messages():
    secret = "SENSITIVE-CUSTOMER-DESCRIPTION"
    sheet = _sheet([["DESCRIPTION", "QUANTITY"], [secret, "SECRET-INVALID-QUANTITY"]])
    result = _extract(sheet)
    _review(result)
    assert result.candidates
    assert all(secret not in issue.message and "SECRET-INVALID-QUANTITY" not in issue.message for issue in result.issues)


def test_optional_terminal_colon_cannot_silently_overwrite_another_field_alias():
    data = load_packing_extraction_config().model_dump(mode="python")
    data["items"]["unit"].append("QTY:")
    config = PackingExtractionConfig.model_validate(data)
    try:
        extractor = DeterministicPackingItemExtractor(config)
    except ValueError:
        return  # Rejecting an extractor-specific normalized collision is safe.
    result = extractor.extract(_book(_sheet([
        ["DESCRIPTION", "QTY", "GROSS WEIGHT"], ["Synthetic Gear", "2", "12.50"],
    ])))
    _review(result)
    assert all(item.unit is None or item.unit.normalized_value != "2" for item in result.items)
