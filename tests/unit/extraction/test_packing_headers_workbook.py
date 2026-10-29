"""Synthetic workbook evidence and adversarial safety tests for TASK-007B."""

from datetime import date, datetime, time
from decimal import Decimal

from openpyxl.utils.cell import get_column_letter
import pytest

from customs_ai.domain.enums import ValueOrigin
from customs_ai.extraction.enums import ExtractionStatus
from customs_ai.extraction.packing_headers import DeterministicPackingHeaderExtractor
from customs_ai.parsers.models import ParsedCell, ParsedSheet, ParsedWorkbook


def cell(row, column, value, **kwargs):
    return ParsedCell(
        coordinate=f"{get_column_letter(column)}{row}", row=row, column=column,
        raw_value=value, data_type=kwargs.pop("data_type", "s"), **kwargs,
    )


def sheet(cells, name="Synthetic Packing", **kwargs):
    return ParsedSheet(name=name, index=0, state=kwargs.pop("state", "visible"), cells=cells, **kwargs)


def workbook(*sheets, file_type="xlsx"):
    return ParsedWorkbook(
        document_id="DOC-PACKING-WORKBOOK", file_type=file_type,
        parser_used="openpyxl" if file_type == "xlsx" else "xlrd",
        workbook_name="synthetic.xlsx", sheet_count=len(sheets),
        sheets=[tab.model_copy(update={"index": index}) for index, tab in enumerate(sheets)],
    )


def extract(*sheets):
    return DeterministicPackingHeaderExtractor().extract(workbook(*sheets))


def value(result, path):
    node = result.packing_list
    for segment in path.split("."):
        node = getattr(node, segment)
    return node


def assert_review(result):
    assert result.status == ExtractionStatus.NEEDS_REVIEW
    assert any(issue.is_review_required for issue in result.issues)
    assert result.packing_list.items == []


def test_all_sixteen_header_paths_at_arbitrary_workbook_locations():
    examples = [
        ("PACKING LIST NO", "000045", "packing_list_number", "000045"),
        ("PACKING LIST DATE", "2026/10/29", "date", date(2026, 10, 29)),
        ("SELLER", "Synthetic Seller", "seller.name", "Synthetic Seller"),
        ("SELLER TAX ID", "000012", "seller.tax_id", "000012"),
        ("SELLER ADDRESS", "Seller Road 1", "seller.address", "Seller Road 1"),
        ("BUYER", "Synthetic Buyer", "buyer.name", "Synthetic Buyer"),
        ("BUYER TAX ID", "000013", "buyer.tax_id", "000013"),
        ("BUYER ADDRESS", "Buyer Road 2", "buyer.address", "Buyer Road 2"),
        ("SHIPMENT REFERENCE", "000077", "shipment_reference", "000077"),
        ("TOTAL PACKAGES", "0", "package_count", Decimal("0")),
        ("PACKAGE TYPE", "Carton", "package_type", "Carton"),
        ("TOTAL GROSS WEIGHT", "12345678901234567890.1200", "gross_weight", Decimal("12345678901234567890.1200")),
        ("TOTAL NET WEIGHT", "15.2500", "net_weight", Decimal("15.2500")),
        ("DIMENSIONS", "10 x 20 x 30 CM", "dimensions", "10 x 20 x 30 CM"),
        ("TOTAL VOLUME", "0.20", "volume", Decimal("0.20")),
        ("INVOICE NO", "000006", "invoice_references", "000006"),
    ]
    cells = [part for row, (label, raw, _, _) in enumerate(examples, start=17)
             for part in (cell(row, 7, label), cell(row, 8, raw))]
    result = extract(sheet(cells))
    assert result.status == ExtractionStatus.EXTRACTED
    for row, (_, raw, path, expected) in enumerate(examples, start=17):
        field = value(result, path)
        if path == "invoice_references":
            assert len(field) == 1
            field = field[0]
        assert field.normalized_value == expected
        assert field.raw_value == raw
        assert field.source_document_id == "DOC-PACKING-WORKBOOK"
        assert field.sheet == "Synthetic Packing" and field.cell == f"H{row}"
        assert field.page is None and field.origin == ValueOrigin.EXTRACTED
        assert field.extraction_method and 0 < field.confidence <= 1
    assert result.packing_list.gross_weight.normalized_value.as_tuple() == Decimal("12345678901234567890.1200").as_tuple()
    assert result.packing_list.items == []


@pytest.mark.parametrize("labels", [
    ("SELLER", "BUYER", "TOTAL PACKAGES", "GROSS WEIGHT", "NET WEIGHT"),
    ("NGƯỜI BÁN", "NGƯỜI MUA", "TỔNG SỐ KIỆN", "TRỌNG LƯỢNG CẢ BÌ", "TRỌNG LƯỢNG TỊNH"),
    ("卖方", "买方", "总件数", "毛重", "净重"),
])
def test_explicit_multilingual_roles_counts_and_weights(labels):
    cells = [part for row, (label, raw) in enumerate(zip(labels, ["Seller A", "Buyer B", "3", "25.50", "20.50"]), start=10)
             for part in (cell(row, 4, label), cell(row, 5, raw))]
    result = extract(sheet(cells))
    assert result.status == ExtractionStatus.EXTRACTED
    assert result.packing_list.seller.name.normalized_value == "Seller A"
    assert result.packing_list.buyer.name.normalized_value == "Buyer B"
    assert result.packing_list.package_count.normalized_value == Decimal("3")
    assert result.packing_list.gross_weight.normalized_value == Decimal("25.50")
    assert result.packing_list.net_weight.normalized_value == Decimal("20.50")


@pytest.mark.parametrize("foreign_role", ["SHIPPER", "EXPORTER", "CONSIGNEE", "IMPORTER", "NGƯỜI GỬI HÀNG", "NGƯỜI NHẬP KHẨU", "发货人", "收货人", "出口商", "进口商"])
def test_shipping_roles_do_not_establish_seller_or_buyer(foreign_role):
    result = extract(sheet([cell(1, 1, "PACKING LIST NO: 000001"), cell(2, 1, foreign_role), cell(2, 2, "Synthetic Company")]))
    assert result.packing_list.packing_list_number.normalized_value == "000001"
    assert result.packing_list.seller is None and result.packing_list.buyer is None
    assert all(candidate.field_path not in {"seller.name", "buyer.name"} for candidate in result.candidates)


def test_longest_party_labels_are_not_party_names():
    result = extract(sheet([
        cell(1, 1, "SELLER ADDRESS: Seller Address"), cell(2, 1, "BUYER ADDRESS: Buyer Address"),
        cell(3, 1, "SELLER TAX ID: 000035"), cell(4, 1, "BUYER VAT NO: 000036"),
    ]))
    assert result.status == ExtractionStatus.EXTRACTED
    assert result.packing_list.seller.name is None and result.packing_list.buyer.name is None
    assert result.packing_list.seller.address.normalized_value == "Seller Address"
    assert result.packing_list.buyer.address.normalized_value == "Buyer Address"
    assert result.packing_list.seller.tax_id.normalized_value == "000035"
    assert result.packing_list.buyer.tax_id.normalized_value == "000036"


@pytest.mark.parametrize("label", ["NO", "DATE", "TOTAL", "QTY", "WEIGHT", "SHIPMENT DATE", "DELIVERY DATE", "PACKING LIST NOTATION", "SELLER ADDRESSING", "BUYERX", "ITEM GROSS WEIGHT"])
def test_unknown_and_false_prefix_labels_never_establish_a_header(label):
    result = extract(sheet([cell(1, 1, label), cell(1, 2, "000001")]))
    assert_review(result)
    assert result.candidates == []
    assert not any(field is not None for name, field in result.packing_list if name not in {"items", "invoice_references"})


@pytest.mark.parametrize("prose", [
    "Description of goods: Seller: Synthetic Company",
    "Synthetic pump with PACKING LIST NO: 000001",
    "NOTES: TOTAL PACKAGES: 10",
    "Product  GROSS WEIGHT: 20.50",
])
def test_label_like_substrings_inside_prose_are_not_headers(prose):
    result = extract(sheet([cell(7, 9, prose)]))
    assert_review(result)
    assert result.candidates == []


def test_packing_and_invoice_numbers_and_multiple_visible_sheets():
    result = extract(
        sheet([cell(7, 5, "PACKING LIST NO"), cell(7, 6, "000011"), cell(8, 5, "INVOICE NO"), cell(8, 6, "000012")], "First"),
        sheet([cell(12, 9, "INVOICE REF.: 000013"), cell(13, 9, "INVOICE NO.: 000012")], "Second"),
    )
    assert result.status == ExtractionStatus.EXTRACTED
    assert result.packing_list.packing_list_number.normalized_value == "000011"
    refs = result.packing_list.invoice_references
    assert {field.normalized_value for field in refs} == {"000012", "000013"}
    assert len(refs) == 2
    assert next(field for field in refs if field.normalized_value == "000013").sheet == "Second"
    assert next(field for field in refs if field.normalized_value == "000013").cell == "I12"
    assert len([candidate for candidate in result.candidates if candidate.field_path == "invoice_references"]) == 3


@pytest.mark.parametrize("combined", ["000011, 000012", "000011; 000012", "000011 / 000012", "000011 and 000012", "000011\n000012"])
def test_ambiguous_combined_invoice_references_are_not_guessed(combined):
    result = extract(sheet([cell(1, 1, "INVOICE REFERENCES"), cell(1, 2, combined)]))
    assert_review(result)
    assert result.packing_list.invoice_references == []
    assert any(candidate.field_path == "invoice_references" and candidate.raw_value == combined for candidate in result.candidates)


def test_raw_cell_evidence_is_not_replaced_by_displayed_value():
    result = extract(sheet([cell(1, 1, "PACKING LIST NO"), cell(1, 2, " 00AbC ", displayed_value="FAKE")]))
    field = result.packing_list.packing_list_number
    assert field.raw_value == " 00AbC " and field.normalized_value == "00AbC"
    assert result.candidates[0].raw_value == " 00AbC "


@pytest.mark.parametrize("raw", [0, 0.0, "0.00"])
def test_numeric_zero_is_valid_evidence(raw):
    result = extract(sheet([cell(11, 4, "TOTAL PACKAGES"), cell(11, 5, raw)]))
    assert result.status == ExtractionStatus.EXTRACTED
    assert result.packing_list.package_count.normalized_value == Decimal("0")
    assert result.packing_list.package_count.cell == "E11"


@pytest.mark.parametrize("raw", ["1,200", "1.200", "KG 100", "100 KG", True, time(12, 30), float("inf")])
def test_unsafe_numeric_values_remain_reviewable(raw):
    result = extract(sheet([cell(1, 1, "GROSS WEIGHT"), cell(1, 2, raw)]))
    assert_review(result)
    assert result.packing_list.gross_weight is None
    assert any(candidate.field_path == "gross_weight" for candidate in result.candidates)


@pytest.mark.parametrize("raw", ["01/02/2026", "2026-02-30", True, time(12, 30), 45000])
def test_ambiguous_or_wrong_type_dates_are_not_coerced(raw):
    result = extract(sheet([cell(1, 1, "PACKING LIST DATE"), cell(1, 2, raw)]))
    assert_review(result)
    assert result.packing_list.date is None


@pytest.mark.parametrize("raw", [date(2026, 10, 29), datetime(2026, 10, 29, 13, 45)])
def test_typed_excel_dates_keep_date_contract(raw):
    result = extract(sheet([cell(1, 1, "PACKING LIST DATE"), cell(1, 2, raw)]))
    assert result.status == ExtractionStatus.EXTRACTED
    assert result.packing_list.date.normalized_value == date(2026, 10, 29)
    assert type(result.packing_list.date.normalized_value) is date
    assert result.packing_list.date.raw_value == str(raw)


@pytest.mark.parametrize("metadata", [{"hidden_rows": [2]}, {"hidden_columns": ["B"]}, {"state": "hidden"}, {"state": "veryHidden"}])
def test_hidden_business_content_cannot_be_trusted(metadata):
    result = extract(sheet([cell(2, 1, "TOTAL PACKAGES"), cell(2, 2, 55)], **metadata))
    assert_review(result)
    assert result.packing_list.package_count is None


def test_hidden_conflicting_content_does_not_leave_unqualified_success():
    result = extract(
        sheet([cell(1, 1, "PACKING LIST NO: 000001")], "Visible"),
        sheet([cell(1, 1, "PACKING LIST NO: 000002")], "Hidden", state="hidden"),
    )
    assert_review(result)
    assert result.packing_list.packing_list_number is None or result.packing_list.packing_list_number.normalized_value == "000001"
    assert all(candidate.raw_value != "000002" for candidate in result.candidates) or result.status == ExtractionStatus.NEEDS_REVIEW


@pytest.mark.parametrize("formula_metadata", [{"is_formula": True}, {"data_type": "f"}, {}])
def test_formula_value_is_opaque_even_with_a_cached_display(formula_metadata):
    result = extract(sheet([cell(1, 1, "TOTAL PACKAGES"), cell(1, 2, "=1+2", displayed_value=3, **formula_metadata)]))
    assert_review(result)
    assert result.packing_list.package_count is None
    assert any(candidate.raw_value == "=1+2" for candidate in result.candidates)


@pytest.mark.parametrize("formula_metadata", [{"is_formula": True}, {"data_type": "f"}, {}])
def test_formula_cannot_manufacture_label(formula_metadata):
    result = extract(sheet([cell(1, 1, '=CONCAT("PACKING LIST NO", ": 000001")', displayed_value="PACKING LIST NO: 000001", **formula_metadata)]))
    assert_review(result)
    assert result.packing_list.packing_list_number is None
    assert result.candidates == []


def test_inline_formula_value_never_produces_business_data():
    result = extract(sheet([cell(1, 1, "PACKING LIST NO: =CONCAT(0, 1)")]))
    assert_review(result)
    assert result.packing_list.packing_list_number is None
    assert result.candidates[0].raw_value == "=CONCAT(0, 1)"


def test_merged_label_and_actual_value_coordinate_are_preserved():
    result = extract(sheet([cell(7, 3, "PACKING LIST NO"), cell(7, 5, "000015")], merged_ranges=["C7:D7"]))
    assert result.status == ExtractionStatus.EXTRACTED
    assert result.packing_list.packing_list_number.normalized_value == "000015"
    assert result.packing_list.packing_list_number.cell == "E7"


def test_merged_value_uses_actual_anchor_and_does_not_create_phantom_cells():
    result = extract(sheet([cell(7, 3, "PACKING LIST NO"), cell(7, 4, "000015")], merged_ranges=["D7:F7"]))
    assert result.status == ExtractionStatus.EXTRACTED
    assert result.packing_list.packing_list_number.cell == "D7"
    assert all(candidate.cell != "E7" and candidate.cell != "F7" for candidate in result.candidates)


def test_empty_merged_region_cannot_manufacture_a_value():
    result = extract(sheet([cell(7, 3, "PACKING LIST NO"), cell(7, 5, None)], merged_ranges=["D7:F7"]))
    assert_review(result)
    assert result.packing_list.packing_list_number is None
    assert not result.candidates


def test_unrelated_nonadjacent_value_is_not_selected_by_offset_search():
    result = extract(sheet([cell(7, 3, "PACKING LIST NO"), cell(7, 6, "Unrelated"), cell(8, 3, "Unrelated Below")]))
    assert_review(result)
    assert result.packing_list.packing_list_number is None


def test_item_column_headers_are_not_document_weights_or_package_count():
    result = extract(sheet([
        cell(2, 1, "PACKING LIST NO: 000010"),
        cell(8, 1, "ITEM CODE"), cell(8, 2, "QUANTITY"), cell(8, 3, "GROSS WEIGHT"), cell(8, 4, "NET WEIGHT"),
        cell(9, 1, "000001"), cell(9, 2, 50), cell(9, 3, "15.50"), cell(9, 4, "10.50"),
        cell(12, 1, "TOTAL GROSS WEIGHT: 45.50"), cell(13, 1, "TOTAL NET WEIGHT: 40.50"),
    ]))
    assert result.packing_list.packing_list_number.normalized_value == "000010"
    assert result.packing_list.package_count is None
    assert result.packing_list.gross_weight.normalized_value == Decimal("45.50")
    assert result.packing_list.net_weight.normalized_value == Decimal("40.50")
    assert result.packing_list.items == []
    assert not any(candidate.cell in {"C9", "D9"} for candidate in result.candidates)


def test_dimensions_stay_text_and_never_calculate_volume():
    result = extract(sheet([cell(1, 1, "DIMENSIONS"), cell(1, 2, "10 x 20 x 30 CM")]))
    assert result.status == ExtractionStatus.EXTRACTED
    assert result.packing_list.dimensions.normalized_value == "10 x 20 x 30 CM"
    assert result.packing_list.volume is None


@pytest.mark.parametrize("numbers", [["000001", "000002"], ["000001", "000002", "000001"], ["000001", "000002", "000003"]])
def test_scalar_conflicts_are_sticky_without_majority_or_first_last_selection(numbers):
    result = extract(sheet([cell(index, 1, "PACKING LIST NO" + ": " + number) for index, number in enumerate(numbers, start=1)]))
    assert_review(result)
    assert result.packing_list.packing_list_number is None
    assert len([candidate for candidate in result.candidates if candidate.field_path == "packing_list_number"]) == len(numbers)


def test_repeated_equivalent_numeric_values_are_not_a_conflict():
    result = extract(sheet([cell(1, 1, "TOTAL PACKAGES: 2"), cell(2, 1, "TOTAL PACKAGES: 2.00")]))
    assert result.status == ExtractionStatus.EXTRACTED
    assert result.packing_list.package_count.normalized_value == Decimal("2")
    assert len([candidate for candidate in result.candidates if candidate.field_path == "package_count"]) == 2


def test_unparseable_candidate_alongside_valid_candidate_stays_unresolved():
    result = extract(sheet([cell(1, 1, "GROSS WEIGHT: 2.50"), cell(2, 1, "GROSS WEIGHT: 1,200")]))
    assert_review(result)
    assert result.packing_list.gross_weight is None
    assert len([candidate for candidate in result.candidates if candidate.field_path == "gross_weight"]) == 2


@pytest.mark.parametrize("file_type", ["xlsx", "xls"])
def test_both_approved_workbook_formats_are_supported(file_type):
    result = DeterministicPackingHeaderExtractor().extract(workbook(sheet([cell(1, 1, "P/L NO: 000021")]), file_type=file_type))
    assert result.status == ExtractionStatus.EXTRACTED
    assert result.packing_list.packing_list_number.normalized_value == "000021"


def test_empty_workbook_returns_explicit_review_without_fabrication():
    result = extract(sheet([]))
    assert_review(result)
    assert result.packing_list.invoice_references == []
    assert result.candidates == []


def test_duplicate_sheet_indices_cannot_misidentify_provenance():
    document = workbook(
        sheet([cell(1, 1, "PACKING LIST NO: 000001")], "First"),
        sheet([cell(1, 1, "PACKING LIST NO: 000002")], "Second"),
    )
    document.sheets[1].index = 0
    result = DeterministicPackingHeaderExtractor().extract(document)
    assert_review(result)
    assert result.packing_list.packing_list_number is None
    assert result.candidates == []


def test_inconsistent_row_column_and_coordinate_cannot_ground_a_field():
    inconsistent = cell(1, 2, "000001").model_copy(update={"coordinate": "C1"})
    result = extract(sheet([cell(1, 1, "PACKING LIST NO"), inconsistent]))
    assert_review(result)
    assert result.packing_list.packing_list_number is None
    assert result.candidates == []


def test_populated_merged_interior_cannot_manufacture_source_value():
    result = extract(sheet([cell(7, 3, "PACKING LIST NO"), cell(7, 4, "000001")], merged_ranges=["C7:D7"]))
    assert_review(result)
    assert result.packing_list.packing_list_number is None


def test_partially_hidden_merged_label_is_not_a_visible_label():
    result = extract(sheet([cell(7, 3, "PACKING LIST NO"), cell(7, 5, "000001")], merged_ranges=["C7:D7"], hidden_columns=["D"]))
    assert_review(result)
    assert result.packing_list.packing_list_number is None


@pytest.mark.parametrize("metadata", [{"is_formula": True}, {"data_type": "f"}])
def test_cached_formula_value_without_expression_is_still_unsafe(metadata):
    result = extract(sheet([cell(1, 1, "TOTAL PACKAGES"), cell(1, 2, 3, displayed_value=3, **metadata)]))
    assert_review(result)
    assert result.packing_list.package_count is None


@pytest.mark.parametrize("metadata", [{"is_formula": True}, {"data_type": "f"}])
def test_formula_flag_cannot_turn_cached_label_text_into_evidence(metadata):
    result = extract(sheet([cell(1, 1, "PACKING LIST NO", **metadata), cell(1, 2, "000001")]))
    assert_review(result)
    assert result.packing_list.packing_list_number is None
    assert result.candidates == []


def test_issue_messages_do_not_contain_source_business_values():
    secret = "SYNTHETIC-CONFIDENTIAL-12345"
    result = extract(sheet([cell(1, 1, "GROSS WEIGHT: " + secret)]))
    assert_review(result)
    assert secret not in " ".join(issue.message for issue in result.issues)
    assert result.candidates[0].raw_value == secret


def test_real_xlsx_parser_to_extractor_retains_source_and_merged_provenance(tmp_path):
    from hashlib import sha256
    from openpyxl import Workbook
    from customs_ai.parsers.excel import XlsxParser

    path = tmp_path / "synthetic-packing.xlsx"
    original = Workbook()
    ws = original.active
    ws.title = "Synthetic PL"
    ws.merge_cells("G17:H17")
    ws["G17"] = "PACKING LIST NO"
    ws["I17"] = "000017"
    ws["G18"] = "INVOICE NO"
    ws["H18"] = "000018"
    ws["G19"] = "TOTAL PACKAGES"
    ws["H19"] = 0
    original.save(path)
    original.close()
    before = sha256(path.read_bytes()).hexdigest()

    result = DeterministicPackingHeaderExtractor().extract(XlsxParser().parse("DOC-REAL-XLSX", path))

    assert result.status == ExtractionStatus.EXTRACTED
    assert result.packing_list.packing_list_number.normalized_value == "000017"
    assert result.packing_list.packing_list_number.cell == "I17"
    assert result.packing_list.packing_list_number.sheet == "Synthetic PL"
    assert result.packing_list.invoice_references[0].cell == "H18"
    assert result.packing_list.package_count.normalized_value == Decimal("0")
    assert sha256(path.read_bytes()).hexdigest() == before


def test_real_xls_parser_to_extractor_retains_source_and_leading_zeros(tmp_path):
    from hashlib import sha256
    import xlwt
    from customs_ai.parsers.excel import XlsParser

    path = tmp_path / "synthetic-packing.xls"
    original = xlwt.Workbook()
    ws = original.add_sheet("Legacy PL")
    ws.write(16, 6, "PACKING LIST NO")
    ws.write(16, 7, "000019")
    ws.write(17, 6, "TOTAL GROSS WEIGHT")
    ws.write(17, 7, "25.5000")
    original.save(str(path))
    before = sha256(path.read_bytes()).hexdigest()

    result = DeterministicPackingHeaderExtractor().extract(XlsParser().parse("DOC-REAL-XLS", path))

    assert result.status == ExtractionStatus.EXTRACTED
    assert result.packing_list.packing_list_number.normalized_value == "000019"
    assert result.packing_list.packing_list_number.cell == "H17"
    assert result.packing_list.gross_weight.normalized_value.as_tuple() == Decimal("25.5000").as_tuple()
    assert sha256(path.read_bytes()).hexdigest() == before
