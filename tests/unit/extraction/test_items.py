"""Semantic and safety tests for the deterministic Excel item-table boundary."""

from datetime import date, time
from decimal import Decimal
from hashlib import sha256

import openpyxl
import pytest
from openpyxl.utils import get_column_letter
from pydantic import ValidationError

from customs_ai.domain.enums import ValueOrigin
from customs_ai.domain.schema import InvoiceItem
from customs_ai.extraction.enums import ExtractionStatus
from customs_ai.extraction.items import (
    DeterministicInvoiceItemExtractor,
    extract_invoice_items,
)
from customs_ai.parsers.excel import XlsxParser
from customs_ai.parsers.models import ParsedCell, ParsedSheet, ParsedWorkbook


def _cell(row, column, value, **overrides):
    default_type = (
        "f" if overrides.get("is_formula") else
        "b" if isinstance(value, bool) else
        "d" if isinstance(value, (date, time)) else
        "n" if isinstance(value, (int, float)) else "s"
    )
    return ParsedCell(
        coordinate=f"{get_column_letter(column)}{row}",
        row=row,
        column=column,
        raw_value=value,
        data_type=overrides.pop("data_type", default_type),
        **overrides,
    )


def _sheet(rows, *, row=1, column=1, name="Invoice", index=0, **metadata):
    cells = [
        _cell(row + offset, column + cell_offset, value)
        for offset, values in enumerate(rows)
        for cell_offset, value in enumerate(values)
        if value is not None
    ]
    return ParsedSheet(
        name=name,
        index=index,
        state=metadata.pop("state", "visible"),
        cells=cells,
        **metadata,
    )


def _workbook(*sheets):
    return ParsedWorkbook(
        document_id="doc_synthetic_items",
        file_type="xlsx",
        parser_used="openpyxl",
        workbook_name="synthetic.xlsx",
        sheet_count=len(sheets),
        sheets=list(sheets),
    )


def _extract(rows, **sheet_metadata):
    return extract_invoice_items(_workbook(_sheet(rows, **sheet_metadata)))


def _assert_review(result):
    assert result.status == ExtractionStatus.NEEDS_REVIEW
    assert any(issue.is_review_required for issue in result.issues)


def test_discovers_table_at_arbitrary_position_beyond_first_worksheet():
    cover = _sheet([["Synthetic cover"], ["No item table here"]], name="Cover")
    invoice = _sheet(
        [["ITEM CODE", "DESCRIPTION", "QUANTITY", "AMOUNT"], ["00012", "Motor", 2, 50]],
        row=17,
        column=5,
        name="Commercial Invoice",
        index=1,
    )
    result = extract_invoice_items(_workbook(cover, invoice))

    assert result.status == ExtractionStatus.EXTRACTED
    assert len(result.items) == 1
    item = result.items[0]
    assert item.item_code.normalized_value == "00012"
    assert item.quantity.normalized_value == Decimal("2")
    assert item.item_code.sheet == "Commercial Invoice"
    assert item.item_code.cell == "E18"
    assert item.amount.cell == "H18"


@pytest.mark.parametrize(
    "labels",
    [
        [" description: ", " UNIT\nPRICE", "QUANTITY"],
        ["MÔ TẢ HÀNG HÓA", "ĐƠN GIÁ", "SỐ LƯỢNG"],
        ["货物名称", "单价", "数量"],
    ],
)
def test_whole_cell_aliases_cover_english_vietnamese_chinese(labels):
    result = _extract([labels, ["Synthetic motor", "12,5", 2]])

    assert result.status == ExtractionStatus.EXTRACTED
    assert len(result.items) == 1
    assert result.items[0].description.normalized_value == "Synthetic motor"
    assert result.items[0].unit_price.normalized_value == Decimal("12.5")


def test_all_fourteen_canonical_fields_preserve_raw_and_provenance():
    labels = [
        "NO", "ITEM CODE", "PART NUMBER", "MODEL", "DESCRIPTION", "BRAND",
        "MANUFACTURER", "QUANTITY", "UNIT", "UNIT PRICE", "AMOUNT", "ORIGIN",
        "WEIGHT", "PACKAGE INFORMATION",
    ]
    values = [
        "0001", "000012", "P-001", "M-0002", "  acme  Motor  ", "Brand-X",
        "Factory-Y", 2, "PCS", "12.50", "25.00", "VN", "1.25", "2 cartons",
    ]
    fields = [
        "item_number", "item_code", "part_number", "model", "description", "brand",
        "manufacturer", "quantity", "unit", "unit_price", "amount", "origin",
        "weight", "package_information",
    ]
    result = _extract([labels, values], row=7, column=3)

    assert result.status == ExtractionStatus.EXTRACTED
    assert isinstance(result.items[0], InvoiceItem)
    for offset, (name, raw) in enumerate(zip(fields, values)):
        field = getattr(result.items[0], name)
        assert field.raw_value == str(raw)
        assert field.origin == ValueOrigin.EXTRACTED
        assert field.source_document_id == "doc_synthetic_items"
        assert field.sheet == "Invoice"
        assert field.cell == f"{get_column_letter(3 + offset)}8"
        assert field.page is None
    assert result.items[0].item_code.normalized_value == "000012"
    assert result.items[0].description.normalized_value == "acme  Motor"
    assert result.items[0].weight.normalized_value == Decimal("1.25")


def test_native_and_string_numeric_values_and_zero_are_not_falsy_missing():
    result = _extract(
        [
            ["DESCRIPTION", "QUANTITY", "UNIT PRICE", "AMOUNT"],
            ["Free sample", 0.0, 0, "0.00"],
            ["US formatted", 2, 1200.5, "2,401.00"],
            ["EU formatted", "12,5", "1.200,50", "15006.25"],
        ]
    )

    assert result.status == ExtractionStatus.EXTRACTED
    assert len(result.items) == 3
    assert result.items[0].quantity.normalized_value == 0
    assert result.items[0].unit_price.normalized_value == 0
    assert result.items[0].amount.normalized_value == Decimal("0.00")
    assert result.items[1].unit_price.normalized_value == Decimal("1200.5")
    assert result.items[2].unit_price.normalized_value == Decimal("1200.50")


def test_unresolved_numeric_values_keep_partial_rows_and_evidence():
    for invalid in ["1,200", "1.200", "1,200,000", "USD 12", True, time(14, 30), date(2026, 1, 1), float("nan"), float("inf")]:
        result = _extract([["DESCRIPTION", "QUANTITY", "AMOUNT"], ["Motor", invalid, 10]])

        _assert_review(result)
        assert len(result.items) == 1
        assert result.items[0].description.normalized_value == "Motor"
        assert result.items[0].quantity is None
        assert result.items[0].amount.normalized_value == Decimal("10")
        assert any(candidate.cell == "B2" for candidate in result.candidates)


def test_formula_values_are_unresolved_even_when_displayed_cache_is_supplied():
    sheet = _sheet([["DESCRIPTION", "QUANTITY", "UNIT PRICE", "AMOUNT"], ["Motor", 2, 5]])
    sheet.cells.append(
        _cell(2, 4, '=IF(TRUE,"  AMOUNT: 999",0)', data_type="f", is_formula=True, displayed_value=999)
    )
    result = extract_invoice_items(_workbook(sheet))

    _assert_review(result)
    assert len(result.items) == 1
    assert result.items[0].amount is None
    assert result.items[0].quantity.normalized_value == Decimal("2")
    formula = next(candidate for candidate in result.candidates if candidate.cell == "D2")
    assert formula.raw_value == '=IF(TRUE,"  AMOUNT: 999",0)'
    assert formula.source_document_id == "doc_synthetic_items"
    assert formula.sheet == "Invoice"


def test_formula_cells_and_expression_text_cannot_supply_table_headers():
    for formula_flag in [False, True]:
        sheet = _sheet([[None, "QTY"], ["Motor", 2]])
        sheet.cells.append(
            _cell(1, 1, '  ="DESCRIPTION"', is_formula=formula_flag, displayed_value="DESCRIPTION")
        )
        result = extract_invoice_items(_workbook(sheet))
        _assert_review(result)
        assert result.items == []


def test_identifier_formula_cannot_invent_an_item_code():
    sheet = _sheet([["DESCRIPTION", "ITEM CODE", "QTY"], ["Motor", None, 2]])
    expression = '="  ITEM CODE: INVENTED | QTY: 999"'
    sheet.cells.append(_cell(2, 2, expression, is_formula=True, data_type="f"))
    result = extract_invoice_items(_workbook(sheet))

    _assert_review(result)
    assert len(result.items) == 1
    assert result.items[0].item_code is None
    assert result.items[0].quantity.normalized_value == Decimal("2")
    assert any(candidate.raw_value == expression for candidate in result.candidates)


def test_missing_fields_stay_none_and_incomplete_relevant_rows_are_retained():
    result = _extract(
        [["DESCRIPTION", "QTY", "AMOUNT"], ["Complete motor", 2, 10], ["Missing quantity", None, 5], [None, 1, 3]]
    )

    _assert_review(result)
    assert len(result.items) == 3
    assert result.items[1].description.normalized_value == "Missing quantity"
    assert result.items[1].quantity is None
    assert result.items[2].description is None
    assert result.items[2].quantity.normalized_value == Decimal("1")
    assert all(item.origin is None and item.unit_price is None for item in result.items)


def test_blank_business_row_terminates_even_when_unrelated_note_column_is_populated():
    result = _extract(
        [
            ["DESCRIPTION", "QTY", "AMOUNT", "REMARKS"],
            ["Real item", 2, 10],
            [None, None, None, "Do not continue table based on this note"],
            ["Unrelated numeric information", 90, 900],
        ]
    )

    assert len(result.items) == 1
    assert result.items[0].description.normalized_value == "Real item"
    assert not any(candidate.cell in {"A4", "B4", "C4"} for candidate in result.candidates)


def test_implicit_blank_row_gap_terminates_the_table():
    sheet = _sheet([["DESCRIPTION", "QTY"], ["Real item", 1]])
    sheet.cells.extend([_cell(5, 1, "Other data"), _cell(5, 2, 99)])
    result = extract_invoice_items(_workbook(sheet))

    assert len(result.items) == 1
    assert result.items[0].description.normalized_value == "Real item"


def test_exact_total_summary_rows_never_become_items_or_restart_on_footer_data():
    for marker in ["TOTAL", "GRAND TOTAL", "TỔNG CỘNG", "合计"]:
        result = _extract(
            [["DESCRIPTION", "QTY", "AMOUNT"], ["Motor", 2, 10], [marker, None, 10], ["Footer data", 99, 999]]
        )
        assert len(result.items) == 1
        assert result.items[0].description.normalized_value == "Motor"


def test_product_names_starting_with_total_are_preserved():
    result = _extract(
        [["DESCRIPTION", "QTY", "AMOUNT"], ["TOTALIZER", 1, 10], ["TOTAL CONTROL", 2, 20], ["TOTAL LOSS", 3, 30]]
    )

    assert result.status == ExtractionStatus.EXTRACTED
    assert [item.description.normalized_value for item in result.items] == ["TOTALIZER", "TOTAL CONTROL", "TOTAL LOSS"]


def test_exact_total_with_real_item_identity_is_retained_with_review():
    result = _extract([["ITEM CODE", "DESCRIPTION", "QTY", "UNIT PRICE", "AMOUNT"], ["SKU-001", "TOTAL", 1, 10, 10]])

    _assert_review(result)
    assert len(result.items) == 1
    assert result.items[0].item_code.normalized_value == "SKU-001"
    assert result.items[0].description.raw_value == "TOTAL"


def test_subtotal_boundary_is_conservative_about_unextracted_sections():
    result = _extract([["DESCRIPTION", "QTY", "AMOUNT"], ["Motor", 1, 10], ["SUBTOTAL", None, 10], ["Another section", 2, 20]])

    _assert_review(result)
    assert len(result.items) == 1
    assert result.items[0].description.normalized_value == "Motor"


def test_material_hidden_rows_are_excluded_and_surface_review():
    result = _extract(
        [["DESCRIPTION", "QTY"], ["Visible motor", 1], ["Hidden draft", 99], ["Visible cable", 2]],
        hidden_rows=[3],
    )

    _assert_review(result)
    assert [item.description.normalized_value for item in result.items] == ["Visible motor", "Visible cable"]


def test_hidden_and_very_hidden_tables_are_never_merged_with_visible_items():
    for state in ["hidden", "veryHidden"]:
        visible = _sheet([["DESCRIPTION", "QTY"], ["Visible motor", 1]])
        hidden = _sheet([["DESCRIPTION", "QTY"], ["Hidden draft", 99]], name="Draft", index=1, state=state)
        result = extract_invoice_items(_workbook(visible, hidden))

        _assert_review(result)
        assert all(item.description.normalized_value != "Hidden draft" for item in result.items)
        assert len(result.items) <= 1


def test_hidden_header_does_not_create_a_trusted_table():
    result = _extract([["DESCRIPTION", "QTY"], ["Motor", 1]], hidden_rows=[1])

    _assert_review(result)
    assert result.items == []


def test_duplicate_alias_columns_are_ambiguous_even_when_data_agrees():
    result = _extract([["DESCRIPTION", "QTY", "QUANTITY"], ["Motor", 2, 2]])

    _assert_review(result)
    assert result.items == []


def test_merged_header_spanning_columns_is_not_assigned_arbitrarily():
    result = _extract([["DESCRIPTION", None, "QTY"], ["Motor", None, 2]], merged_ranges=["A1:B1"])

    _assert_review(result)
    assert result.items == []


def test_two_stacked_credible_tables_require_review_without_arbitrary_selection():
    result = _extract([["DESCRIPTION", "QTY"], ["Motor", 1], [], ["DESCRIPTION", "QTY"], ["Cable", 2]])

    _assert_review(result)
    assert result.items == []


def test_two_side_by_side_tables_do_not_collapse_duplicate_columns():
    result = _extract([["DESCRIPTION", "QTY", None, "DESCRIPTION", "QTY"], ["Motor", 1, None, "Cable", 2]])

    _assert_review(result)
    assert result.items == []


def test_tables_across_visible_sheets_are_not_blindly_appended():
    first = _sheet([["DESCRIPTION", "QTY"], ["Motor", 1]], name="Invoice A")
    second = _sheet([["DESCRIPTION", "QTY"], ["Cable", 2]], name="Invoice B", index=1)
    result = extract_invoice_items(_workbook(first, second))

    _assert_review(result)
    assert result.items == []


def test_contiguous_repeated_identical_header_is_a_structural_continuation():
    result = _extract([["DESCRIPTION", "QTY"], ["Motor", 1], ["DESCRIPTION", "QTY"], ["Cable", 2]])

    assert result.status == ExtractionStatus.EXTRACTED
    assert [item.description.normalized_value for item in result.items] == ["Motor", "Cable"]
    assert [item.description.cell for item in result.items] == ["A2", "A4"]


def test_isolated_generic_labels_and_no_quantity_do_not_prove_product_identity():
    for rows in [
        [["NO", "QTY"], [1, 3]],
        [["AMOUNT"], [100]],
        [["CODE"], ["SKU-1"]],
        [["DESCRIPTION NOTES", "QUANTITY"], ["Not an exact description header", 3]],
    ]:
        result = _extract(rows)
        _assert_review(result)
        assert result.items == []


def test_notes_footer_in_business_column_does_not_become_a_product():
    result = _extract([["DESCRIPTION", "QTY"], ["Motor", 2], ["REMARKS: synthetic footer"], ["Unrelated information", 999]])

    _assert_review(result)
    assert len(result.items) == 1
    assert result.items[0].description.normalized_value == "Motor"


def test_real_xlsx_parser_integration_keeps_raw_evidence_and_source_immutable(tmp_path):
    path = tmp_path / "synthetic_invoice.xlsx"
    workbook = openpyxl.Workbook()
    cover = workbook.active
    cover.title = "Cover"
    cover["A1"] = "Synthetic Commercial Invoice"
    sheet = workbook.create_sheet("Items")
    for column, value in enumerate(["ITEM CODE", "DESCRIPTION", "QTY", "UNIT PRICE", "AMOUNT"], start=4):
        sheet.cell(13, column, value)
    for column, value in enumerate(["000123", "  Synthetic Motor  ", 0, 12.5, "=F14*G14"], start=4):
        sheet.cell(14, column, value)
    workbook.save(path)
    workbook.close()
    before = sha256(path.read_bytes()).hexdigest()

    parsed = XlsxParser().parse("doc_synthetic_items", path)
    result = DeterministicInvoiceItemExtractor().extract(parsed)

    _assert_review(result)
    assert len(result.items) == 1
    item = result.items[0]
    assert item.item_code.normalized_value == "000123"
    assert item.description.raw_value == "  Synthetic Motor  "
    assert item.quantity.normalized_value == Decimal("0")
    assert item.unit_price.normalized_value == Decimal("12.5")
    assert item.amount is None
    assert item.item_code.sheet == "Items"
    assert item.item_code.cell == "D14"
    assert any(candidate.cell == "H14" and candidate.raw_value == "=F14*G14" for candidate in result.candidates)
    assert sha256(path.read_bytes()).hexdigest() == before


def test_result_is_typed_strict_serializable_and_source_grounded():
    result = _extract([["DESCRIPTION", "QTY"], ["Motor", 2]])

    restored = type(result).model_validate_json(result.model_dump_json())
    assert restored.items[0].quantity.normalized_value == Decimal("2")
    with pytest.raises(ValidationError):
        type(result).model_validate({**result.model_dump(), "undeclared": True})
    for item in result.items:
        for name in InvoiceItem.model_fields:
            field = getattr(item, name)
            if field is not None:
                assert field.origin == ValueOrigin.EXTRACTED
                assert field.raw_value is not None
                assert field.source_document_id == "doc_synthetic_items"


def test_side_by_side_disjoint_field_groups_cannot_create_a_hybrid_item():
    result = _extract(
        [["ITEM CODE", "QTY", None, "DESCRIPTION", "AMOUNT"],
         ["SKU-A", 2, None, "Different cable", 999]]
    )
    _assert_review(result)
    assert result.items == []
    assert {candidate.cell for candidate in result.candidates} >= {"A1", "B1", "D1", "E1"}


def test_hidden_summary_terminates_before_footer_numeric_data():
    result = _extract(
        [["DESCRIPTION", "QTY", "AMOUNT"], ["Motor", 2, 10],
         ["TOTAL", None, 10], ["Unrelated footer data", 99, 999]],
        hidden_rows=[3],
    )
    _assert_review(result)
    assert [item.description.normalized_value for item in result.items] == ["Motor"]
    assert not any(candidate.cell in {"A4", "B4", "C4"} for candidate in result.candidates)


def test_note_marker_with_unrelated_numbers_is_not_a_product():
    result = _extract(
        [["DESCRIPTION", "QTY"], ["Motor", 2], ["REMARKS: shipment note", 999]]
    )
    _assert_review(result)
    assert [item.description.normalized_value for item in result.items] == ["Motor"]
    assert any(candidate.raw_value == "REMARKS: shipment note" for candidate in result.candidates)


def test_formatted_native_identifier_preserves_raw_and_display_zero_padding():
    sheet = _sheet([["ITEM CODE", "QTY"], [12, 1]])
    sheet.cells = [
        cell.model_copy(update={"number_format": "000000"}) if cell.coordinate == "A2" else cell
        for cell in sheet.cells
    ]
    result = extract_invoice_items(_workbook(sheet))
    assert result.status == ExtractionStatus.EXTRACTED
    assert result.items[0].item_code.raw_value == "12"
    assert result.items[0].item_code.normalized_value == "000012"


def test_inconsistent_cell_provenance_is_rejected_without_fabricating_coordinates():
    sheet = _sheet([["DESCRIPTION", "QTY"], ["Motor", 1]])
    sheet.cells = [
        cell.model_copy(update={"coordinate": "Z99"}) if cell.coordinate == "B2" else cell
        for cell in sheet.cells
    ]
    result = extract_invoice_items(_workbook(sheet))
    _assert_review(result)
    assert result.items == []


def test_unrelated_numbers_do_not_establish_product_description_identity():
    result = _extract([["DESCRIPTION", "QTY"], [123, 456]])
    _assert_review(result)
    assert result.items == []
    assert {candidate.raw_value for candidate in result.candidates} == {"123", "456"}


def test_footer_labels_support_fullwidth_colon_and_separate_following_text():
    for footer in ["备注：付款说明", "GHI CHÚ: thanh toán", "NOTES synthetic footer"]:
        result = _extract([["DESCRIPTION", "QTY"], ["Motor", 1], [footer, 999]])
        _assert_review(result)
        assert [item.description.normalized_value for item in result.items] == ["Motor"]
        assert any(candidate.raw_value == footer for candidate in result.candidates)


def test_compound_or_hidden_column_mapping_requires_review():
    for sheet in [
        _sheet([["DESCRIPTION", "QTY / AMOUNT"], ["Motor", 2]]),
        _sheet([["DESCRIPTION", "QTY"], ["Motor", 2]], hidden_columns=["B"]),
    ]:
        result = extract_invoice_items(_workbook(sheet))
        _assert_review(result)
        assert result.items == []


def test_formula_only_identity_does_not_self_certify_a_table():
    sheet = _sheet([["DESCRIPTION", "QTY"], [None, 2]])
    sheet.cells.append(_cell(2, 1, '="Motor"', is_formula=True))
    result = extract_invoice_items(_workbook(sheet))
    _assert_review(result)
    assert result.items == []
    assert any(candidate.raw_value == '="Motor"' for candidate in result.candidates)


def test_invalid_input_is_reported_without_source_or_service_side_effects():
    unsupported = extract_invoice_items(object())
    assert unsupported.status == ExtractionStatus.FAILED
    assert unsupported.items == []
    workbook = _workbook(_sheet([["DESCRIPTION", "QTY"], ["Motor", 1]]))
    missing_source = extract_invoice_items(workbook.model_copy(update={"document_id": " "}))
    assert missing_source.status == ExtractionStatus.FAILED
    assert missing_source.items == []


def test_real_legacy_xls_numeric_cells_and_source_are_preserved(tmp_path):
    import xlwt
    from customs_ai.parsers.excel import XlsParser

    path = tmp_path / "synthetic_items.xls"
    book = xlwt.Workbook()
    sheet = book.add_sheet("Items")
    for column, value in enumerate(["ITEM CODE", "QTY", "AMOUNT"]):
        sheet.write(12, column + 4, value)
    for column, value in enumerate(["000123", 0, 12.5]):
        sheet.write(13, column + 4, value)
    book.save(str(path))
    before = sha256(path.read_bytes()).hexdigest()
    parsed = XlsParser().parse("doc_synthetic_items", path)
    result = extract_invoice_items(parsed)
    assert result.status == ExtractionStatus.EXTRACTED
    assert result.items[0].item_code.normalized_value == "000123"
    assert result.items[0].quantity.normalized_value == Decimal("0.0")
    assert result.items[0].amount.normalized_value == Decimal("12.5")
    assert result.items[0].item_code.cell == "E14"
    assert sha256(path.read_bytes()).hexdigest() == before


def test_approved_total_amount_labels_are_summary_boundaries_not_products():
    for marker in ["TOTAL AMOUNT", "TOTAL INVOICE VALUE", "TỔNG TIỀN", "TOTAL PRICE"]:
        result = _extract(
            [["DESCRIPTION", "QTY", "AMOUNT"], ["Motor", 2, 10], [marker, None, 10]]
        )
        assert result.status == ExtractionStatus.EXTRACTED
        assert [item.description.normalized_value for item in result.items] == ["Motor"]

