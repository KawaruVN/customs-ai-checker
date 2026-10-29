"""Structural table discovery and row semantics for deterministic Packing items."""

from decimal import Decimal

from openpyxl.utils.cell import get_column_letter
import pytest

from customs_ai.domain.enums import ValueOrigin
from customs_ai.domain.schema import PackingItem
from customs_ai.extraction.enums import ExtractionStatus
from customs_ai.extraction.packing_items import (
    DeterministicPackingItemExtractor,
    extract_packing_items,
)
from customs_ai.extraction.packing_models import PackingItemExtractionResult
from customs_ai.parsers.models import ParsedCell, ParsedSheet, ParsedWorkbook


def _cell(row, column, raw):
    return ParsedCell(
        coordinate=f"{get_column_letter(column)}{row}", row=row, column=column,
        raw_value=raw, data_type="n" if isinstance(raw, (int, float)) else "s",
    )


def _sheet(rows, *, row=1, column=1, name="Synthetic Packing", index=0):
    return ParsedSheet(
        name=name, index=index, state="visible",
        cells=[_cell(row + offset, column + column_offset, raw)
               for offset, values in enumerate(rows)
               for column_offset, raw in enumerate(values) if raw is not None],
    )


def _book(*sheets):
    return ParsedWorkbook(
        document_id="DOC-PACKING-ITEM-TABLES", file_type="xlsx", parser_used="openpyxl",
        workbook_name="synthetic.xlsx", sheet_count=len(sheets), sheets=list(sheets),
    )


def _extract(rows, **metadata):
    return extract_packing_items(_book(_sheet(rows, **metadata)))


def _review(result):
    assert isinstance(result, PackingItemExtractionResult)
    assert result.status == ExtractionStatus.NEEDS_REVIEW
    assert result.issues and any(issue.is_review_required for issue in result.issues)


@pytest.mark.parametrize("labels", [
    ["ITEM CODE", "DESCRIPTION", "QUANTITY", "UNIT", "PACKAGE NO.", "GROSS WEIGHT", "NET WEIGHT"],
    ["MÃ HÀNG", "MÔ TẢ HÀNG HÓA", "SỐ LƯỢNG", "ĐƠN VỊ TÍNH", "SỐ HIỆU KIỆN", "TRỌNG LƯỢNG CẢ BÌ", "TRỌNG LƯỢNG TỊNH"],
    ["货号", "品名", "数量", "单位", "箱号", "毛重", "净重"],
], ids=["english", "vietnamese", "chinese"])
def test_all_seven_fields_multilingual_with_exact_raw_and_source_paths(labels):
    source = ["000012", "MiXeD Synthetic Motor", "0.00", "PCS", "000045",
              "12345678901234567890.1200", "12.5000"]
    result = _extract([labels, source], row=8, column=5)
    assert result.status == ExtractionStatus.EXTRACTED and not result.issues
    assert len(result.items) == 1 and isinstance(result.items[0], PackingItem)
    fields = ["item_code", "description", "quantity", "unit", "package_number", "gross_weight", "net_weight"]
    expected = [source[0], source[1], Decimal(source[2]), source[3], source[4], Decimal(source[5]), Decimal(source[6])]
    for column, path, raw, normalized in zip(range(5, 12), fields, source, expected):
        field = getattr(result.items[0], path)
        assert field.raw_value == raw and field.normalized_value == normalized
        assert field.origin == ValueOrigin.EXTRACTED
        assert field.source_document_id == "DOC-PACKING-ITEM-TABLES"
        assert field.sheet == "Synthetic Packing" and field.cell == f"{get_column_letter(column)}9"
        assert field.page is None and field.extraction_method
        assert field.confidence is not None and 0 < field.confidence <= 1
        candidate = next(candidate for candidate in result.candidates if candidate.cell == field.cell)
        assert candidate.raw_value == raw
        assert candidate.field_path == f"tables[0:8].rows[9].columns[{column}].{path}"
    assert result.items[0].quantity.normalized_value.as_tuple() == Decimal("0.00").as_tuple()
    assert result.items[0].gross_weight.normalized_value.as_tuple() == Decimal(source[5]).as_tuple()
    assert result.items[0].net_weight.normalized_value.as_tuple() == Decimal("12.5000").as_tuple()


@pytest.mark.parametrize("start_row,start_column", [(1, 1), (17, 5), (104, 30), (9, 100)])
def test_table_can_start_at_arbitrary_row_and_column(start_row, start_column):
    result = _extract([["ITEM CODE", "QUANTITY"], ["000012", "2"]], row=start_row, column=start_column)
    assert result.status == ExtractionStatus.EXTRACTED
    assert result.items[0].item_code.normalized_value == "000012"
    assert result.items[0].quantity.cell == f"{get_column_letter(start_column + 1)}{start_row + 1}"


@pytest.mark.parametrize("order", [
    ["QUANTITY", "DESCRIPTION", "ITEM CODE"],
    ["DESCRIPTION", "ITEM CODE", "QUANTITY"],
    ["ITEM CODE", "QUANTITY", "DESCRIPTION"],
])
def test_column_order_is_discovered_from_labels(order):
    values = {"ITEM CODE": "000012", "DESCRIPTION": "Synthetic Pump", "QUANTITY": "2.50"}
    result = _extract([order, [values[label] for label in order]])
    assert result.status == ExtractionStatus.EXTRACTED
    assert result.items[0].item_code.normalized_value == "000012"
    assert result.items[0].description.normalized_value == "Synthetic Pump"
    assert result.items[0].quantity.normalized_value == Decimal("2.50")


def test_extra_unrelated_columns_do_not_become_canonical_fields():
    result = _extract([
        ["INTERNAL NOTE", "ITEM CODE", "QUANTITY", "CUSTOMER COMMENT"],
        ["Unrelated left", "000012", "2", "Unrelated right"],
    ])
    assert result.status == ExtractionStatus.EXTRACTED
    assert result.items[0].item_code.cell == "B2" and result.items[0].quantity.cell == "C2"
    assert all(candidate.cell not in {"A2", "D2"} for candidate in result.candidates)
    assert set(PackingItem.model_fields) == {"item_code", "description", "quantity", "unit", "package_number", "gross_weight", "net_weight"}


def test_visible_table_after_cover_sheet_uses_actual_sheet_evidence():
    cover = _sheet([["Synthetic cover"], ["No goods table"]], name="Cover", index=0)
    goods = _sheet([["DESCRIPTION", "GROSS WEIGHT"], ["Synthetic Motor", "12.50"]],
                   row=19, column=4, name="Goods", index=1)
    result = extract_packing_items(_book(cover, goods))
    assert result.status == ExtractionStatus.EXTRACTED
    assert result.items[0].description.sheet == "Goods" and result.items[0].description.cell == "D20"
    assert result.items[0].gross_weight.cell == "E20"
    assert result.items[0].quantity is None


@pytest.mark.parametrize("numeric_label,path", [
    ("QUANTITY", "quantity"), ("GROSS WEIGHT", "gross_weight"), ("NET WEIGHT", "net_weight"),
])
def test_identity_plus_any_meaningful_item_quantity_or_weight_is_supported(numeric_label, path):
    result = _extract([["ITEM CODE", numeric_label], ["000012", "0.00"]])
    assert result.status == ExtractionStatus.EXTRACTED
    assert getattr(result.items[0], path).normalized_value == Decimal("0.00")
    assert result.items[0].package_number is None


@pytest.mark.parametrize("rows", [
    [["PACKAGE NO.", "QUANTITY"], ["000045", "10"]],
    [["GROSS WEIGHT", "NET WEIGHT"], ["12.50", "11.50"]],
    [["TOTAL PACKAGES", "GROSS WEIGHT"], ["10", "12.50"]],
    [["QUANTITY", "UNIT"], ["10", "PCS"]],
    [["DESCRIPTION", "QUANTITY"], ["123456", "10"]],
    [["ITEM CODE", "QUANTITY"], [12, "10"]],
    [["DESCRIPTION", "QUANTITY"], ["Synthetic Motor", "ambiguous"]],
    [["ITEM CODE", "QUANTITY"]],
    [["ITEM CODE EXTRA", "QUANTITY"], ["000012", "10"]],
    [["NOTES ABOUT ITEM CODE", "QUANTITY"], ["000012", "10"]],
    [["DESCRIPTION: ITEM CODE / QUANTITY"], ["000012 / 10"]],
    [["ITEM CODE", "TOTAL PACKAGES"], ["000012", "10"]],
    [["ITEM CODE", "TOTAL GROSS WEIGHT"], ["000012", "10"]],
])
def test_false_regions_do_not_certify_a_goods_table(rows):
    result = _extract(rows)
    _review(result)
    assert result.items == []


def test_identical_contiguous_repeated_header_is_one_supported_continuation():
    header = ["ITEM CODE", "DESCRIPTION", "QUANTITY"]
    result = _extract([header, ["000012", "Motor", "2"], header, ["000013", "Pump", "3"]])
    assert result.status == ExtractionStatus.EXTRACTED
    assert [item.item_code.normalized_value for item in result.items] == ["000012", "000013"]
    assert [item.quantity.cell for item in result.items] == ["C2", "C4"]
    assert all(candidate.cell not in {"A3", "B3", "C3"} for candidate in result.candidates
               if ".rows[" in candidate.field_path)


@pytest.mark.parametrize("marker", ["TOTAL", "GRAND TOTAL", "TỔNG CỘNG", "合计", "NOTES", "REMARKS", "GHI CHÚ", "备注"])
def test_exact_summary_or_footer_terminates_without_reading_later_unheaded_rows(marker):
    result = _extract([
        ["DESCRIPTION", "QUANTITY"], ["Motor", "2"], [marker, "2"], ["Unrelated later row", "99"],
    ])
    assert len(result.items) == 1 and result.items[0].description.normalized_value == "Motor"
    assert all(item.description is None or item.description.normalized_value != marker for item in result.items)
    assert not any(candidate.cell in {"A4", "B4"} for candidate in result.candidates)


@pytest.mark.parametrize("marker", ["SUBTOTAL", "SUB TOTAL", "小计"])
def test_subtotal_ends_region_with_review_instead_of_automatic_section_concatenation(marker):
    result = _extract([["DESCRIPTION", "QUANTITY"], ["Motor", "2"], [marker, "2"], ["Second section", "99"]])
    _review(result)
    assert len(result.items) == 1 and result.items[0].description.normalized_value == "Motor"


@pytest.mark.parametrize("description", [
    "TOTALIZER", "TOTAL CONTROL", "TOTAL LOSS SENSOR", "SUMMARY CONTROL UNIT",
    "NOTES DISPLAY MODULE", "Motor with TOTAL markings", "Pump with NOTES label",
])
def test_real_product_description_merely_containing_summary_words_is_preserved(description):
    result = _extract([["DESCRIPTION", "QUANTITY"], [description, "2"]])
    assert result.status == ExtractionStatus.EXTRACTED
    assert len(result.items) == 1 and result.items[0].description.normalized_value == description


@pytest.mark.parametrize("gap", [[], [None, None], [None, None, "Unmapped note cannot bridge mapped gap"]])
def test_blank_mapped_row_or_implicit_gap_terminates_table(gap):
    result = _extract([["ITEM CODE", "QUANTITY"], ["000012", "2"], gap, ["000013", "3"]])
    assert len(result.items) == 1 and result.items[0].item_code.normalized_value == "000012"
    assert all(candidate.cell not in {"A4", "B4"} for candidate in result.candidates)


def test_same_header_after_blank_is_competing_table_not_automatic_continuation():
    header = ["ITEM CODE", "QUANTITY"]
    result = _extract([header, ["000012", "2"], [], header, ["000013", "3"]])
    _review(result)
    assert result.items == []
    assert {"A2", "B2", "A5", "B5"}.issubset({candidate.cell for candidate in result.candidates})


def test_different_header_immediately_after_goods_establishes_competing_region():
    result = _extract([
        ["ITEM CODE", "QUANTITY"], ["000012", "2"],
        ["DESCRIPTION", "NET WEIGHT"], ["Synthetic Pump", "3.50"],
    ])
    _review(result)
    assert result.items == []
    assert {"000012", "Synthetic Pump"}.issubset({candidate.raw_value for candidate in result.candidates})


@pytest.mark.parametrize("count", [2, 3])
@pytest.mark.parametrize("separate_sheets", [False, True])
def test_multiple_credible_tables_never_select_first_or_majority(count, separate_sheets):
    if separate_sheets:
        sheets = [_sheet([["ITEM CODE", "QUANTITY"], [f"CODE-{index}", "2"]], name=f"Goods {index}", index=index)
                  for index in range(count)]
        document = _book(*sheets)
    else:
        rows = []
        for index in range(count):
            rows.extend([["ITEM CODE", "QUANTITY"], [f"CODE-{index}", "2"], []])
        document = _book(_sheet(rows))
    result = extract_packing_items(document)
    _review(result)
    assert result.items == []
    assert {f"CODE-{index}" for index in range(count)}.issubset({candidate.raw_value for candidate in result.candidates})


@pytest.mark.parametrize("other_rows", [
    [["DESCRIPTION", "QUANTITY"], ["12345", "99"]],
    [["ITEM CODE", "QUANTITY"]],
    [["PACKAGE NO.", "QUANTITY"], ["000045", "99"]],
])
def test_noncredible_region_does_not_outvote_supported_table(other_rows):
    result = _extract([["ITEM CODE", "QUANTITY"], ["000012", "2"], [], *other_rows])
    assert len(result.items) == 1 and result.items[0].item_code.normalized_value == "000012"
    # Recognized but ungrounded product tables must be reviewable; pure unrelated
    # package metadata may be ignored rather than manufacturing another table.
    if other_rows[0][0] in {"DESCRIPTION", "ITEM CODE"}:
        _review(result)


def test_one_supported_table_plus_unrelated_notes_does_not_create_second_table():
    result = _extract([
        ["NOTES: columns are descriptive"], [],
        ["ITEM CODE", "DESCRIPTION", "QUANTITY"], ["000012", "Synthetic Motor", "2"],
        ["NOTES", "Other text", None],
    ])
    assert len(result.items) == 1
    assert result.items[0].item_code.normalized_value == "000012"


def test_partial_rows_in_supported_region_preserve_safe_fields_without_shifting():
    result = _extract([
        ["ITEM CODE", "DESCRIPTION", "QUANTITY"],
        ["000012", "Complete Motor", "2"],
        ["000013", None, None],
        [None, None, "3"],
    ])
    _review(result)
    assert len(result.items) == 3
    assert result.items[0].quantity.normalized_value == Decimal("2")
    assert result.items[1].item_code.normalized_value == "000013"
    assert result.items[1].quantity is None
    assert result.items[2].item_code is None and result.items[2].description is None
    assert result.items[2].quantity.normalized_value == Decimal("3")
    assert result.items[2].quantity.cell == "C4"


def test_description_continuation_is_not_silently_joined_into_previous_item():
    result = _extract([
        ["ITEM CODE", "DESCRIPTION", "QUANTITY"],
        ["000012", "Synthetic Motor", "2"],
        [None, "Additional physical-row description", None],
    ])
    _review(result)
    assert result.items[0].description.normalized_value == "Synthetic Motor"
    assert result.items[0].description.raw_value == "Synthetic Motor"
    assert any(candidate.cell == "B3" and candidate.raw_value == "Additional physical-row description"
               for candidate in result.candidates)


def test_missing_package_and_weights_are_optional_when_product_quantity_is_supported():
    result = _extract([["ITEM CODE", "QUANTITY"], ["000012", "2"]])
    assert result.status == ExtractionStatus.EXTRACTED
    assert result.items[0].package_number is None
    assert result.items[0].gross_weight is None and result.items[0].net_weight is None


def test_two_items_may_share_package_identifier_without_deduplicating_business_rows():
    result = _extract([
        ["ITEM CODE", "QUANTITY", "PACKAGE NO."],
        ["000012", "2", "000045"], ["000013", "3", "000045"],
    ])
    assert result.status == ExtractionStatus.EXTRACTED
    assert len(result.items) == 2
    assert [item.package_number.normalized_value for item in result.items] == ["000045", "000045"]
    assert [item.package_number.cell for item in result.items] == ["C2", "C3"]


@pytest.mark.parametrize("combined", ["000045,000046", "000045;000046", "000045\n000046"])
def test_multiple_package_identifiers_in_one_cell_are_reviewed_without_guessing(combined):
    result = _extract([["ITEM CODE", "QUANTITY", "PACKAGE NO."], ["000012", "2", combined]])
    _review(result)
    assert result.items[0].package_number is None
    assert any(candidate.raw_value == combined and candidate.cell == "C2" for candidate in result.candidates)


def test_candidates_physical_row_identity_does_not_use_canonical_item_index():
    result = _extract([["ITEM CODE", "QUANTITY"], ["000012", "2"], ["000013", "3"]], row=19, column=7)
    assert result.status == ExtractionStatus.EXTRACTED
    assert len(result.items) == 2
    paths = {candidate.field_path for candidate in result.candidates}
    assert "tables[0:19].rows[20].columns[7].item_code" in paths
    assert "tables[0:19].rows[21].columns[8].quantity" in paths
    assert all(not path.startswith("items[20]") and not path.startswith("items[21]") for path in paths)


def test_extractor_does_not_mutate_parsed_workbook_or_leak_private_values():
    document = _book(_sheet([
        ["ITEM CODE", "DESCRIPTION", "QUANTITY"],
        ["PRIVATE-000012", "Private synthetic description", "1,200"],
        ["000013", "Supported Item", "2"],
    ]))
    before = document.model_dump(mode="python")
    result = DeterministicPackingItemExtractor().extract(document)
    _review(result)
    assert document.model_dump(mode="python") == before
    assert all("PRIVATE-000012" not in issue.message and "Private synthetic description" not in issue.message
               and "1,200" not in issue.message for issue in result.issues)
