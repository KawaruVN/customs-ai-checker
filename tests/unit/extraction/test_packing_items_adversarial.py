"""Development probes for false table credibility and scalar package ambiguity."""

import pytest
from openpyxl.utils.cell import get_column_letter

from customs_ai.extraction.enums import ExtractionStatus
from customs_ai.extraction.packing_items import DeterministicPackingItemExtractor
from customs_ai.parsers.models import ParsedCell, ParsedSheet, ParsedWorkbook


def workbook(rows, *, start_row=1, start_column=1):
    cells = [
        ParsedCell(
            coordinate=f"{get_column_letter(start_column + column)}{start_row + row}",
            row=start_row + row, column=start_column + column,
            raw_value=value, data_type="s" if isinstance(value, str) else "n",
        )
        for row, values in enumerate(rows)
        for column, value in enumerate(values)
        if value is not None
    ]
    return ParsedWorkbook(
        document_id="PACK-ADVERSARIAL", file_type="xlsx", parser_used="openpyxl",
        workbook_name="synthetic.xlsx", sheet_count=1,
        sheets=[ParsedSheet(name="Synthetic", index=0, state="visible", cells=cells)],
    )


@pytest.mark.parametrize("combined", [
    "0001, 0002", "0001; 0002", "0001 / 0002", "0001 and 0002",
    "0001\n0002", "0001|0002",
])
def test_multiple_package_identifiers_do_not_become_one_clean_scalar(combined):
    result = DeterministicPackingItemExtractor().extract(workbook([
        ["DESCRIPTION", "QTY", "PACKAGE NO."],
        ["Synthetic Pump", "2", combined],
    ]))
    assert result.status == ExtractionStatus.NEEDS_REVIEW
    assert all(item.package_number is None for item in result.items)
    assert any(candidate.raw_value == combined for candidate in result.candidates)
    assert any(issue.is_review_required for issue in result.issues)


def test_split_identity_and_quantity_rows_cannot_certify_product_table():
    result = DeterministicPackingItemExtractor().extract(workbook([
        ["DESCRIPTION", "QTY"],
        ["Synthetic Pump", None],
        [None, "2"],
    ]))
    assert result.status == ExtractionStatus.NEEDS_REVIEW
    assert result.items == []
    assert {candidate.raw_value for candidate in result.candidates} >= {"Synthetic Pump", "2"}


@pytest.mark.parametrize("label", ["TOTAL GROSS WEIGHT", "TOTAL NET WEIGHT", "TOTAL PACKAGES"])
def test_document_total_labels_do_not_supply_item_quantitative_column(label):
    result = DeterministicPackingItemExtractor().extract(workbook([
        ["DESCRIPTION", label],
        ["Synthetic Pump", "2"],
    ]))
    assert result.status == ExtractionStatus.NEEDS_REVIEW
    assert result.items == []


@pytest.mark.parametrize("identity_header", [
    "NOTES: DESCRIPTION", "Product DESCRIPTION", "DESCRIPTION OF GOODS EXAMPLE",
    "品名说明", "GHI CHÚ: MÔ TẢ HÀNG HÓA",
])
def test_alias_mentioned_inside_notes_or_false_prefix_does_not_establish_identity(identity_header):
    result = DeterministicPackingItemExtractor().extract(workbook([
        [identity_header, "QTY"],
        ["Synthetic Pump", "2"],
    ]))
    assert result.status == ExtractionStatus.NEEDS_REVIEW
    assert result.items == []


def test_explicit_note_context_on_header_row_does_not_promote_example_to_goods():
    result = DeterministicPackingItemExtractor().extract(workbook([
        ["NOTES: illustrative example", "DESCRIPTION", "QTY"],
        [None, "Example product", "2"],
    ]))
    assert result.status == ExtractionStatus.NEEDS_REVIEW
    assert result.items == []


def test_candidate_path_uses_physical_row_identity_without_inventing_item_index():
    result = DeterministicPackingItemExtractor().extract(workbook([
        ["DESCRIPTION", "QTY"],
        ["Synthetic Pump", "2"],
    ], start_row=23, start_column=28))
    assert result.status == ExtractionStatus.EXTRACTED
    assert len(result.items) == 1
    assert result.items[0].description.cell == "AB24"
    body_candidates = [candidate for candidate in result.candidates if not candidate.field_path.startswith("headers.")]
    assert body_candidates and all(".rows[24]." in candidate.field_path for candidate in body_candidates)
    assert all(not candidate.field_path.startswith("items[24]") for candidate in result.candidates)


def test_opaque_single_package_identifier_with_slashes_remains_exact():
    result = DeterministicPackingItemExtractor().extract(workbook([
        ["DESCRIPTION", "QTY", "PACKAGE NO."],
        ["Synthetic Pump", "2", "BOX/2026/0001"],
    ]))
    assert result.status == ExtractionStatus.EXTRACTED
    assert result.items[0].package_number.raw_value == "BOX/2026/0001"
    assert result.items[0].package_number.normalized_value == "BOX/2026/0001"


@pytest.mark.parametrize("placeholder", ["---", "...", "???", "\x00"])
def test_punctuation_or_control_only_code_does_not_certify_goods_identity(placeholder):
    result = DeterministicPackingItemExtractor().extract(workbook([
        ["ITEM CODE", "QTY"],
        [placeholder, "2"],
    ]))
    assert result.status == ExtractionStatus.NEEDS_REVIEW
    assert result.items == []
    assert any(candidate.raw_value == placeholder for candidate in result.candidates)


def test_contiguous_explicit_notes_section_does_not_promote_example_table_to_goods():
    result = DeterministicPackingItemExtractor().extract(workbook([
        ["NOTES: illustrative example"],
        ["DESCRIPTION", "QTY"],
        ["Example product", "2"],
    ]))
    assert result.status == ExtractionStatus.NEEDS_REVIEW
    assert result.items == []


def test_blank_separator_can_establish_table_after_earlier_notes_section():
    result = DeterministicPackingItemExtractor().extract(workbook([
        ["NOTES: illustrative example"],
        [],
        ["DESCRIPTION", "QTY"],
        ["Synthetic Pump", "2"],
    ]))
    assert result.status == ExtractionStatus.EXTRACTED
    assert result.items[0].description.normalized_value == "Synthetic Pump"


def test_hidden_empty_row_cannot_manufacture_visible_separator_after_notes():
    source = workbook([
        ["NOTES: illustrative example"],
        [],
        ["DESCRIPTION", "QTY"],
        ["Example product", "2"],
    ])
    source.sheets[0].hidden_rows = [2]
    result = DeterministicPackingItemExtractor().extract(source)
    assert result.status == ExtractionStatus.NEEDS_REVIEW
    assert result.items == []
