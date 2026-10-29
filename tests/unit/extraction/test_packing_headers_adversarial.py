"""Independent development probes for opaque contexts and explicit role safety."""

import pytest
from decimal import Decimal
from openpyxl.utils.cell import get_column_letter

from customs_ai.extraction.enums import ExtractionStatus
from customs_ai.extraction.packing_headers import DeterministicPackingHeaderExtractor
from customs_ai.parsers.models import ParsedCell, ParsedPdfDocument, ParsedPdfPage, ParsedSheet, ParsedWorkbook
from customs_ai.vision.enums import TextResolutionSource, VerificationLevel
from customs_ai.vision.models import ResolvedPdfDocument, ResolvedPdfPage


def _pdf(text):
    return ParsedPdfDocument(
        document_id="PACK-ADVERSARIAL-PDF", page_count=1, has_text_layer=True,
        needs_ocr=False, pages=[ParsedPdfPage(
            page=1, text=text + "\nSynthetic packing document supporting evidence text.",
        )],
    )


def _workbook(rows):
    cells = [ParsedCell(
        coordinate=f"{get_column_letter(column)}{row}", row=row, column=column,
        raw_value=value, data_type="s",
    ) for row, values in enumerate(rows, start=1)
       for column, value in enumerate(values, start=1) if value is not None]
    return ParsedWorkbook(
        document_id="PACK-ADVERSARIAL-EXCEL", file_type="xlsx", parser_used="openpyxl",
        workbook_name="synthetic.xlsx", sheet_count=1,
        sheets=[ParsedSheet(name="Synthetic", index=0, state="visible", cells=cells)],
    )


@pytest.mark.parametrize("text", [
    "NOTES: Example note  SELLER: Fake Company",
    "DESCRIPTION: Widget  SELLER: Fake Company",
    "REMARKS: Product note | INVOICE NO: FAKE",
    "DESCRIPTION: Widget | TOTAL GROSS WEIGHT: 999",
    "UNKNOWN KEY: Product note\tTOTAL NET WEIGHT: 999",
])
@pytest.mark.parametrize("source", ["pdf", "excel"])
def test_opaque_non_target_values_cannot_promote_embedded_labels(text, source):
    document = _pdf(text) if source == "pdf" else _workbook([[text]])
    result = DeterministicPackingHeaderExtractor().extract(document)
    assert result.status == ExtractionStatus.NEEDS_REVIEW
    assert result.packing_list.seller is None
    assert result.packing_list.invoice_references == []
    assert result.packing_list.gross_weight is None
    assert result.packing_list.net_weight is None
    assert result.candidates == []


@pytest.mark.parametrize("tail", [
    "DESCRIPTION: Widget | SELLER: Fake Company",
    "NOTES: See | INVOICE NO: FAKE",
    "QUANTITY: 2 | GROSS WEIGHT: 999",
])
@pytest.mark.parametrize("source", ["pdf", "excel"])
def test_non_target_boundary_becomes_opaque_after_a_legitimate_header(tail, source):
    text = "PACKING LIST NO: PL1 | " + tail
    document = _pdf(text) if source == "pdf" else _workbook([[text]])
    result = DeterministicPackingHeaderExtractor().extract(document)
    # Retaining the preceding clear header or abstaining on the mixed item row
    # are both safe; a later mention inside opaque value text is never evidence.
    if result.packing_list.packing_list_number is not None:
        assert result.packing_list.packing_list_number.normalized_value == "PL1"
    else:
        assert result.status == ExtractionStatus.NEEDS_REVIEW
    assert result.packing_list.seller is None
    assert result.packing_list.invoice_references == []
    assert result.packing_list.gross_weight is None
    assert all(candidate.field_path == "packing_list_number" for candidate in result.candidates)


@pytest.mark.parametrize("foreign_label", [
    "SHIPPER", "CONSIGNEE", "EXPORTER", "IMPORTER",
    "NGƯỜI GỬI HÀNG", "NGƯỜI NHẬN HÀNG", "发货人", "收货人",
])
@pytest.mark.parametrize("source", ["pdf", "excel"])
def test_bare_foreign_role_label_is_not_a_party_value(foreign_label, source):
    document = (_pdf(f"SELLER:\n{foreign_label}\nSynthetic Shipping Company")
                if source == "pdf" else _workbook([["SELLER", foreign_label]]))
    result = DeterministicPackingHeaderExtractor().extract(document)
    assert result.status == ExtractionStatus.NEEDS_REVIEW
    assert result.packing_list.seller is None
    assert result.packing_list.buyer is None
    assert all(candidate.raw_value != foreign_label for candidate in result.candidates)


def test_product_description_cells_do_not_create_party_headers():
    document = _workbook([
        ["DESCRIPTION", "QUANTITY", "GROSS WEIGHT"],
        ["SELLER: Product Name", 2, "30.50"],
    ])
    result = DeterministicPackingHeaderExtractor().extract(document)
    assert result.status == ExtractionStatus.NEEDS_REVIEW
    assert result.packing_list.seller is None
    assert result.packing_list.gross_weight is None
    assert result.packing_list.items == []
    assert all(candidate.field_path != "seller.name" for candidate in result.candidates)


def test_pdf_product_description_continuation_does_not_create_party_header():
    result = DeterministicPackingHeaderExtractor().extract(_pdf("DESCRIPTION\nSELLER: Product Name"))
    assert result.status == ExtractionStatus.NEEDS_REVIEW
    assert result.packing_list.seller is None
    assert result.packing_list.items == []
    assert all(candidate.field_path != "seller.name" for candidate in result.candidates)


def test_item_row_weight_label_with_inline_value_does_not_establish_document_total():
    document = _workbook([["ITEM CODE", "QUANTITY", "GROSS WEIGHT: 30.50"]])
    result = DeterministicPackingHeaderExtractor().extract(document)
    assert result.status == ExtractionStatus.NEEDS_REVIEW
    assert result.packing_list.gross_weight is None
    assert all(candidate.field_path != "gross_weight" for candidate in result.candidates)


def test_item_context_carries_across_pdf_pages_until_explicit_blank_separator():
    document = _pdf("PACKING LIST NO: PL1\nDESCRIPTION | QTY | GROSS WEIGHT")
    document.page_count = 2
    document.pages.append(ParsedPdfPage(
        page=2, text="SELLER: Product Name\nTOTAL GROSS WEIGHT: 999\n\nBUYER: Real Buyer Company",
    ))
    result = DeterministicPackingHeaderExtractor().extract(document)
    assert result.status == ExtractionStatus.NEEDS_REVIEW
    assert result.packing_list.seller is None and result.packing_list.gross_weight is None
    assert result.packing_list.buyer.name.normalized_value == "Real Buyer Company"
    assert result.packing_list.buyer.name.page == 2


def test_explicit_total_text_inside_excel_description_context_is_still_opaque():
    document = _workbook([
        ["DESCRIPTION", "QUANTITY", "GROSS WEIGHT"],
        ["TOTAL GROSS WEIGHT: 999", 2, "30.50"],
        [],  # An actual blank row separates the subsequent document header.
        ["TOTAL GROSS WEIGHT", "40.50"],
    ])
    result = DeterministicPackingHeaderExtractor().extract(document)
    assert result.status == ExtractionStatus.NEEDS_REVIEW
    assert result.packing_list.gross_weight.normalized_value == Decimal("40.50")
    assert result.packing_list.gross_weight.cell == "B4"
    assert all(candidate.raw_value != "999" for candidate in result.candidates)


def test_multiple_adjacent_label_value_pairs_are_not_mistaken_for_table_columns():
    result = DeterministicPackingHeaderExtractor().extract(_workbook([
        ["PACKING LIST NO", "0001", "TOTAL PACKAGES", "0.00"],
        ["GROSS WEIGHT", "12.50", "NET WEIGHT", "11.50"],
    ]))
    assert result.status == ExtractionStatus.EXTRACTED
    assert result.packing_list.packing_list_number.normalized_value == "0001"
    assert result.packing_list.package_count.normalized_value == Decimal("0.00")
    assert result.packing_list.gross_weight.cell == "B2"
    assert result.packing_list.net_weight.cell == "D2"


def test_omitted_hidden_row_does_not_manufacture_table_context_separator():
    document = _workbook([
        ["DESCRIPTION", "QUANTITY", "GROSS WEIGHT"],
        ["Hidden product", 2, "30.50"],
        ["SELLER: Product Name", 2, "30.50"],
    ])
    document.sheets[0].hidden_rows = [2]
    result = DeterministicPackingHeaderExtractor().extract(document)
    assert result.status == ExtractionStatus.NEEDS_REVIEW
    assert result.packing_list.seller is None


@pytest.mark.parametrize("value", [12, 12.0, True])
def test_numeric_identifier_cannot_silently_lose_excel_display_leading_zeroes(value):
    document = _workbook([["PACKING LIST NO", value]])
    document.sheets[0].cells[1].number_format = "000000"
    result = DeterministicPackingHeaderExtractor().extract(document)
    assert result.status == ExtractionStatus.NEEDS_REVIEW
    assert result.packing_list.packing_list_number is None
    assert result.candidates[0].raw_value == str(value)


@pytest.mark.parametrize("separator", ["\n", "\r", "\r\n"])
@pytest.mark.parametrize("raw", [
    "PACKING LIST NO: PL1{separator}INVOICE NO: INV1",
    "SELLER: Seller Company{separator}BUYER: Buyer Company",
])
def test_multiline_inline_cell_does_not_merge_different_business_fields(raw, separator):
    result = DeterministicPackingHeaderExtractor().extract(_workbook([[raw.format(separator=separator)]]))
    assert result.status == ExtractionStatus.NEEDS_REVIEW
    assert result.packing_list.packing_list_number is None
    assert result.packing_list.invoice_references == []
    assert result.packing_list.seller is None and result.packing_list.buyer is None
    assert len(result.candidates) == 1 and separator in result.candidates[0].raw_value


def test_multiline_adjacent_cell_with_embedded_role_stays_unresolved():
    result = DeterministicPackingHeaderExtractor().extract(_workbook([
        ["SELLER", "Seller Company\nBUYER: Buyer Company"],
    ]))
    assert result.status == ExtractionStatus.NEEDS_REVIEW
    assert result.packing_list.seller is None and result.packing_list.buyer is None


def test_multiline_address_without_another_label_keeps_original_source():
    raw = "Synthetic Road 1\nDistrict 2"
    result = DeterministicPackingHeaderExtractor().extract(_workbook([["SELLER ADDRESS", raw]]))
    assert result.status == ExtractionStatus.EXTRACTED
    assert result.packing_list.seller.address.raw_value == raw
    assert result.packing_list.seller.address.normalized_value == raw
    assert result.packing_list.seller.address.cell == "B1"


@pytest.mark.parametrize("score,accepted", [(0.0, False), (0.1, False), (1.0, True), (None, True)])
def test_supplied_visual_agreement_cannot_contradict_configured_safety_gate(score, accepted):
    document = ResolvedPdfDocument(
        document_id="PACK-VISUAL", page_count=1, requires_review=False, visual_pages=[1],
        pages=[ResolvedPdfPage(
            page=1, text="PACKING LIST NO: PL-0001", source=TextResolutionSource.CROSS_VERIFIED_VISUAL,
            verification_level=VerificationLevel.CROSS_VERIFIED, agreement_score=score,
            confidence=0.95, requires_review=False,
        )],
    )
    result = DeterministicPackingHeaderExtractor().extract(document)
    if accepted:
        assert result.status == ExtractionStatus.EXTRACTED
        assert result.packing_list.packing_list_number.normalized_value == "PL-0001"
    else:
        assert result.status == ExtractionStatus.NEEDS_REVIEW
        assert result.packing_list.packing_list_number is None and result.candidates == []
