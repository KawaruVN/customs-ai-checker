import hashlib
from datetime import datetime

import pytest
from openpyxl import Workbook

from customs_ai.parsers.errors import FileCorruptedError
from customs_ai.parsers.excel import XlsxParser


def _build_xlsx(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice"
    ws["A1"] = "Hóa đơn thử nghiệm"
    ws["B2"] = 12.5
    ws["C3"] = True
    ws["D4"] = datetime(2026, 9, 11, 8, 30)
    ws["E5"] = "=SUM(B2,1)"
    ws["B2"].number_format = "0.00"
    ws.merge_cells("A7:C7")
    ws["A7"] = "Merged"
    ws.row_dimensions[9].hidden = True
    ws.column_dimensions["F"].hidden = True

    hidden = wb.create_sheet("测试")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "中文"
    wb.save(path)
    wb.close()


def _cell(sheet, coordinate):
    return next(cell for cell in sheet.cells if cell.coordinate == coordinate)


def test_xlsx_preserves_cell_and_sheet_provenance(tmp_path):
    path = tmp_path / "synthetic.xlsx"
    _build_xlsx(path)

    parsed = XlsxParser().parse("DOC-XLSX", path)

    assert parsed.file_type == "xlsx"
    assert parsed.parser_used == "openpyxl"
    assert parsed.sheet_count == 2
    assert [sheet.name for sheet in parsed.sheets] == ["Invoice", "测试"]
    assert parsed.sheets[1].state == "hidden"

    invoice = parsed.sheets[0]
    assert _cell(invoice, "A1").raw_value == "Hóa đơn thử nghiệm"
    assert _cell(invoice, "B2").raw_value == 12.5
    assert _cell(invoice, "B2").number_format == "0.00"
    assert _cell(invoice, "C3").raw_value is True
    assert _cell(invoice, "D4").raw_value == datetime(2026, 9, 11, 8, 30)
    formula = _cell(invoice, "E5")
    assert formula.raw_value == "=SUM(B2,1)"
    assert formula.is_formula is True
    assert formula.displayed_value is None
    assert "A7:C7" in invoice.merged_ranges
    assert 9 in invoice.hidden_rows
    assert "F" in invoice.hidden_columns


def test_xlsx_does_not_emit_structurally_empty_cells(tmp_path):
    path = tmp_path / "sparse.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "one"
    ws["XFD1048576"] = None  # inflate dimensions without meaningful content
    wb.save(path)
    wb.close()

    parsed = XlsxParser().parse("DOC-SPARSE", path)
    assert [cell.coordinate for cell in parsed.sheets[0].cells] == ["A1"]


def test_corrupt_xlsx_is_rejected(tmp_path):
    path = tmp_path / "bad.xlsx"
    path.write_bytes(b"PK\x03\x04not-a-real-workbook")

    with pytest.raises(FileCorruptedError) as exc:
        XlsxParser().parse("DOC-BAD", path)
    assert exc.value.code == "FILE_CORRUPTED"


def test_xlsx_source_file_remains_unchanged(tmp_path):
    path = tmp_path / "immutable.xlsx"
    _build_xlsx(path)
    before = hashlib.sha256(path.read_bytes()).hexdigest()

    XlsxParser().parse("DOC-HASH", path)

    assert hashlib.sha256(path.read_bytes()).hexdigest() == before
