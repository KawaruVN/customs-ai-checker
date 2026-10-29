"""Synthetic documents and real production entry points for invoice integration."""

import asyncio
from datetime import date
from hashlib import sha256
from io import BytesIO
from pathlib import Path

from fastapi import UploadFile
import openpyxl
import pytest
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from starlette.datastructures import Headers

from customs_ai.application.composition import build_classification_service, build_invoice_extraction_service
from customs_ai.config import settings
from customs_ai.ingestion.service import IngestionService
from customs_ai.ingestion.validator import SUPPORTED_EXTENSIONS
from customs_ai.repositories.document_extractions import DocumentExtractionRepository
from customs_ai.repositories.source_documents import SourceDocumentRepository


class InvoicePipeline:
    """Fixture generation only; all business processing uses approved modules."""

    def __init__(self, tmp_path):
        self.directory = tmp_path / "synthetic-fixtures"
        self.directory.mkdir()
        self.repo = SourceDocumentRepository()
        self.extractions = DocumentExtractionRepository()

    @staticmethod
    def rows():
        return [
            ["COMMERCIAL INVOICE"],
            ["INVOICE NO", "0000456"],
            ["INVOICE DATE", "2026-10-29"],
            ["SELLER", "Synthetic Exporter"],
            ["BUYER", "Synthetic Importer"],
            ["SHIP TO", "Synthetic Receiver"],
            ["CURRENCY", "USD"],
            ["TOTAL AMOUNT", "123456789012345.1200"],
            ["FREIGHT", "2.50"],
            [],
            [None, None, "ITEM CODE", "DESCRIPTION", "QTY", "UNIT", "UNIT PRICE", "AMOUNT"],
            [None, None, "000009", "Synthetic control motor", "0.00", "PCS", "12.50", "0.00"],
            [None, None, "000010", "TOTAL CONTROL", "2.50", "PCS", "4.00", "10.00"],
            [],
            ["REMARKS", "Synthetic payment terms"],
        ]

    def xlsx(self, rows=None, *, filename="synthetic.xlsx", mutate=None):
        path = self.directory / filename
        book = openpyxl.Workbook()
        book.active.title = "Invoice Data"
        for row in self.rows() if rows is None else rows:
            book.active.append(row)
        if mutate is not None:
            mutate(book)
        book.save(path)
        book.close()
        return path

    def xls(self, rows=None, *, filename="synthetic.xls"):
        import xlwt
        path = self.directory / filename
        book = xlwt.Workbook()
        sheet = book.add_sheet("Invoice Data")
        for r, values in enumerate(self.rows() if rows is None else rows):
            for c, value in enumerate(values):
                if value is not None:
                    if isinstance(value, date):
                        sheet.write(r, c, value, xlwt.easyxf(num_format_str="YYYY-MM-DD"))
                    else:
                        sheet.write(r, c, value)
        book.save(str(path))
        return path

    def pdf(self, pages=None, *, filename="synthetic.pdf"):
        if pages is None:
            pages = [
                "COMMERCIAL INVOICE\nINVOICE NO: 0000456\nINVOICE DATE: 2026-10-29\nSELLER: Synthetic Exporter\nBUYER: Synthetic Importer",
                "TOTAL AMOUNT: 123456789012345.1200\nCURRENCY: USD\nSynthetic invoice page two",
            ]
        path = self.directory / filename
        writer = PdfWriter()
        for text in pages:
            page = writer.add_blank_page(width=612, height=792)
            if text:
                font = DictionaryObject({
                    NameObject("/Type"): NameObject("/Font"),
                    NameObject("/Subtype"): NameObject("/Type1"),
                    NameObject("/BaseFont"): NameObject("/Helvetica"),
                })
                page[NameObject("/Resources")] = DictionaryObject({
                    NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})
                })
                def literal(line):
                    return line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
                commands = " 0 -16 Td ".join(f"({literal(line)}) Tj" for line in text.splitlines())
                stream = DecodedStreamObject()
                stream.set_data(("BT /F1 10 Tf 40 740 Td " + commands + " ET").encode("latin-1"))
                page[NameObject("/Contents")] = writer._add_object(stream)
        with path.open("wb") as handle:
            writer.write(handle)
        return path

    def register(self, path, *, filename=None, shipment_id="SHP-INTEGRATION"):
        upload = UploadFile(
            file=BytesIO(path.read_bytes()), filename=filename or path.name,
            headers=Headers({"content-type": SUPPORTED_EXTENSIONS[path.suffix]}),
        )
        try:
            response, status = asyncio.run(IngestionService(self.repo).ingest(shipment_id, upload))
        finally:
            upload.file.close()
        assert status == 201
        source = self.repo.get_by_document_id(response.document_id)
        assert source.file_hash == sha256(path.read_bytes()).hexdigest()
        assert Path(source.stored_path).resolve().is_relative_to(settings.upload_root.resolve())
        return source

    def classify(self, document_id):
        return build_classification_service(self.repo).classify_document(document_id)

    def extract(self, document_id, *, semantic_provider=None, **preconditions):
        service = build_invoice_extraction_service(self.repo, semantic_provider=semantic_provider)
        return service.extract_document(document_id, **preconditions)

    def source_bytes(self, document_id):
        return Path(self.repo.get_by_document_id(document_id).stored_path).read_bytes()


@pytest.fixture
def invoice_pipeline(tmp_path, monkeypatch):
    # Vision integration tests deliberately replace the factory with controlled
    # local doubles; other tests must never initialize optional OCR engines.
    monkeypatch.setattr(settings.vision, "enabled", False)
    return InvoicePipeline(tmp_path)


class PackingPipeline(InvoicePipeline):
    """Synthetic bytes plus the actual Packing production factories; no extract logic."""

    def __init__(self, tmp_path):
        packing_root = tmp_path / "packing-fixture-root"
        packing_root.mkdir()
        super().__init__(packing_root)
        self.original_hashes = {}

    @staticmethod
    def rows():
        return [
            ["PACKING LIST"],
            ["PACKING LIST NO", "0000456"],
            ["PACKING LIST DATE", "2026-10-29"],
            ["SELLER", "Synthetic Exporter"],
            ["SELLER TAX ID", "00000123"],
            ["SELLER ADDRESS", "Seller Road 1"],
            ["BUYER", "Synthetic Importer"],
            ["BUYER TAX ID", "00000456"],
            ["BUYER ADDRESS", "Buyer Road 2"],
            ["SHIPMENT REFERENCE", "SHIP/2026/0001"],
            ["INVOICE NO", "00000789"],
            ["INVOICE NO", "00000790"],
            ["TOTAL PACKAGES", "0.00"],
            ["PACKAGE TYPE", "Carton"],
            ["TOTAL GROSS WEIGHT", "123456789012345.1200"],
            ["TOTAL NET WEIGHT", "120000000000000.2500"],
            ["DIMENSIONS", "10 x 20 x 30 CM"],
            ["TOTAL VOLUME", "0.20"],
            [],
            [None, None, "ITEM CODE", "DESCRIPTION", "QTY", "UNIT", "PACKAGE NO.", "GROSS WEIGHT", "NET WEIGHT"],
            [None, None, "000009", "Synthetic control motor", "0.00", "PCS", "000001", "3.50", "3.00"],
            [None, None, "000010", "TOTAL CONTROL", "2.50", "PCS", "000002", "12.50", "11.50"],
        ]

    def xlsx(self, rows=None, *, filename="synthetic-packing.xlsx", mutate=None):
        def prepare(book):
            book.active.title = "Packing Data"
            if mutate is not None:
                mutate(book)
        return super().xlsx(self.rows() if rows is None else rows, filename=filename, mutate=prepare)

    def xls(self, rows=None, *, filename="synthetic-packing.xls"):
        import xlwt
        path = self.directory / filename
        book = xlwt.Workbook()
        sheet = book.add_sheet("Packing Data")
        for r, values in enumerate(self.rows() if rows is None else rows):
            for c, value in enumerate(values):
                if value is not None:
                    if isinstance(value, date):
                        sheet.write(r, c, value, xlwt.easyxf(num_format_str="YYYY-MM-DD"))
                    else:
                        sheet.write(r, c, value)
        book.save(str(path))
        return path

    def pdf(self, pages=None, *, filename="synthetic-packing.pdf"):
        if pages is None:
            pages = [
                "PACKING LIST\nPACKING LIST NO: 0000456\nPACKING LIST DATE: 2026-10-29\n"
                "SELLER: Synthetic Exporter\nSELLER TAX ID: 00000123\nSELLER ADDRESS: Seller Road 1\n"
                "BUYER: Synthetic Importer\nBUYER TAX ID: 00000456\nBUYER ADDRESS: Buyer Road 2\n"
                "SHIPMENT REFERENCE: SHIP/2026/0001\nINVOICE NO: 00000789",
                "PACKING LIST\nINVOICE NO: 00000790\nTOTAL PACKAGES: 0.00\nPACKAGE TYPE: Carton\n"
                "TOTAL GROSS WEIGHT: 123456789012345.1200\nTOTAL NET WEIGHT: 120000000000000.2500\n"
                "DIMENSIONS: 10 x 20 x 30 CM\nTOTAL VOLUME: 0.20",
            ]
        return super().pdf(pages, filename=filename)

    def register(self, path, **kwargs):
        original = sha256(path.read_bytes()).hexdigest()
        source = super().register(path, **kwargs)
        assert sha256(path.read_bytes()).hexdigest() == original
        self.original_hashes[source.document_id] = original
        return source

    def extract(self, document_id, **preconditions):
        from customs_ai.application.composition import build_packing_extraction_service
        return build_packing_extraction_service(self.repo).extract_document(document_id, **preconditions)

    def assert_immutable(self, source):
        assert sha256(self.source_bytes(source.document_id)).hexdigest() == self.original_hashes[source.document_id]

    def assert_roundtrip(self, source, outcome):
        from contextlib import closing
        from decimal import Decimal
        import json
        from pydantic import BaseModel
        from customs_ai.domain.schema import CanonicalField
        from customs_ai.repositories.database import get_connection
        from customs_ai.repositories.extraction_models import serialize_result, deserialize_result

        assert outcome.persisted and outcome.record is not None
        fresh = DocumentExtractionRepository()
        by_id = fresh.get_by_extraction_id(outcome.record.extraction_id)
        assert by_id == outcome.record and by_id.result == outcome.result
        assert fresh.get_active(source.document_id) == outcome.active_record
        history = fresh.list_history(source.document_id)
        assert history[-1] == by_id
        assert sum(record.is_active for record in history) <= 1
        assert [record.attempt_number for record in history] == list(range(1, len(history) + 1))
        assert by_id.schema_version == "0.1"
        with closing(get_connection()) as connection:
            payload = connection.execute("SELECT payload_json FROM document_extractions WHERE extraction_id=?",
                                         (by_id.extraction_id,)).fetchone()[0]
        envelope = json.loads(payload)
        assert (envelope["format_version"], envelope["result_type"]) == (2, "packing_list")
        assert deserialize_result(payload) == outcome.result
        assert deserialize_result(serialize_result(by_id.result)) == outcome.result

        def fields(value):
            if isinstance(value, CanonicalField):
                yield value
            elif isinstance(value, BaseModel):
                for _, child in value:
                    yield from fields(child)
            elif isinstance(value, list):
                for child in value:
                    yield from fields(child)
        for original, restored in zip(fields(outcome.result.packing_list), fields(by_id.result.packing_list), strict=True):
            assert original == restored
            assert original.source_document_id == source.document_id
            if isinstance(original.normalized_value, Decimal):
                assert original.normalized_value.as_tuple() == restored.normalized_value.as_tuple()
            assert type(original.normalized_value) is type(restored.normalized_value)
        for original, restored in zip(outcome.result.candidates, by_id.result.candidates, strict=True):
            assert original == restored and type(original.typed_value) is type(restored.typed_value)
            if isinstance(original.typed_value, Decimal):
                assert original.typed_value.as_tuple() == restored.typed_value.as_tuple()
        self.assert_immutable(source)
        return by_id


@pytest.fixture
def packing_pipeline(tmp_path, monkeypatch):
    monkeypatch.setattr(settings.vision, "enabled", False)
    return PackingPipeline(tmp_path)
